"""Editable Word/DOCX exports for native Reports.

The PDF exports are printable snapshots. These DOCX files are intentionally
text/table first so Coordinator can open them in Word, tweak wording, and hand them
off without copying from JSON/CSV.
"""

from __future__ import annotations

import base64
import io
from typing import Any

from docx import Document

from backend_py.domain.calendar_utils import date_range
from backend_py.domain.clinics import build_week_grid, clinic_assignment_rows, week_grid_card_text
from backend_py.domain.exports import build_legend
from backend_py.domain.reports import active_block, detect_conflicts, generate_daily_report, get_rotator

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def build_word_exports(state: dict[str, Any]) -> list[dict[str, str]]:
    block = active_block(state) or {}
    safe_name = _slug(block.get("name") or "schedule")
    files = [
        (f"{safe_name}-schedule.docx", _schedule_docx(state, block)),
        (f"{safe_name}-outpatient-week-grid.docx", _outpatient_week_grid_docx(state, block)),
        (f"{safe_name}-daily-reports.docx", _daily_reports_docx(state, block)),
        (f"{safe_name}-roster-legend.docx", _roster_legend_docx(state, block)),
    ]
    return [
        {
            "name": name,
            "mimeType": DOCX_MIME,
            "base64": base64.b64encode(content).decode("ascii"),
        }
        for name, content in files
    ]


def _schedule_docx(state: dict[str, Any], block: dict[str, Any]) -> bytes:
    document = Document()
    _title(document, block, "Schedule")

    table = document.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    _set_cells(table.rows[0].cells, ["Date", "Inpatient", "Outpatient Service", "Clinic Assignments"])
    clinic_rows = clinic_assignment_rows(state, block)

    for date in _block_dates(block):
        row = table.add_row()
        inpatient = [
            f"{_rotator_label(state, item.get('rotatorId'))} ({item.get('role') or 'Resident'})"
            for item in state.get("inpatientAssignments") or []
            if item.get("date") == date and item.get("role") != "Off"
        ]
        outpatient = [
            f"{item.get('period') or ''} {item.get('clinic') or 'Clinic'}: {_rotator_label(state, item.get('rotatorId'))}"
            for item in state.get("outpatientSessions") or []
            if item.get("date") == date
        ]
        clinics = [_clinic_line(row_) for row_ in clinic_rows if row_.get("date") == date]
        _set_cells(
            row.cells,
            [
                date,
                "\n".join(inpatient) if inpatient else "Open",
                "\n".join(outpatient) if outpatient else "",
                "\n".join(clinics) if clinics else "",
            ],
        )

    document.add_heading("Conflicts", level=2)
    conflicts = detect_conflicts(state)
    if conflicts:
        for conflict in conflicts:
            document.add_paragraph(
                f"{conflict.get('date') or 'No date'} - {conflict.get('severity') or ''}: {conflict.get('title') or ''}"
            )
    else:
        document.add_paragraph("No conflicts detected.")
    return _document_bytes(document)


# Week color bands cycle like the app's Clinics week grid (blue, green,
# orange, purple as light hex shades readable behind black text).
WEEK_BAND_SHADES = ["DCE9F9", "DFF1DF", "FBEAD3", "EBE0F5"]
NO_CLINIC_SHADE = "3B3B3B"

WEEKDAY_COLUMNS = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday"]


def _shade_cell(cell: Any, hex_color: str) -> None:
    # python-docx has no public cell-shading API; set w:shd on tcPr directly.
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    properties = cell._tc.get_or_add_tcPr()
    shade = OxmlElement("w:shd")
    shade.set(qn("w:val"), "clear")
    shade.set(qn("w:fill"), hex_color)
    properties.append(shade)


def _outpatient_week_grid_docx(state: dict[str, Any], block: dict[str, Any]) -> bytes:
    document = Document()
    _title(document, block, "Outpatient Week Grid")

    for week in build_week_grid(state, block):
        document.add_heading(f"{week['label']} ({week['startDate']} to {week['endDate']})", level=2)
        table = document.add_table(rows=3, cols=6)
        table.style = "Table Grid"
        days_by_weekday = {day["weekday"]: day for day in week["days"]}
        band = WEEK_BAND_SHADES[(week["weekIndex"] - 1) % len(WEEK_BAND_SHADES)]

        header = table.rows[0].cells
        _set_cells(header, [""] + [
            f"{weekday[:3].upper()}\n{days_by_weekday[weekday]['date']}" if weekday in days_by_weekday else ""
            for weekday in WEEKDAY_COLUMNS
        ])
        for cell in header:
            _shade_cell(cell, band)

        for row_index, session in enumerate(("AM", "PM"), start=1):
            cells = table.rows[row_index].cells
            _set_cells(cells[:1], [session])
            _shade_cell(cells[0], band)
            for column, weekday in enumerate(WEEKDAY_COLUMNS, start=1):
                day = days_by_weekday.get(weekday)
                if day is None:
                    _set_cells(cells[column:column + 1], [""])
                    continue
                if day.get("holiday"):
                    _set_cells(cells[column:column + 1], [f"{day['holiday']} - No Clinic"])
                    _shade_cell(cells[column], NO_CLINIC_SHADE)
                    continue
                cards = day.get(session) or []
                _set_cells(
                    cells[column:column + 1],
                    ["\n".join(week_grid_card_text(card) for card in cards)],
                )

    return _document_bytes(document)


def _daily_reports_docx(state: dict[str, Any], block: dict[str, Any]) -> bytes:
    document = Document()
    _title(document, block, "Daily Reports")
    for index, date in enumerate(_block_dates(block)):
        if index:
            document.add_page_break()
        document.add_heading(date, level=2)
        for line in generate_daily_report(state, date).splitlines():
            if line.strip():
                document.add_paragraph(line)
            else:
                document.add_paragraph("")
    return _document_bytes(document)


def _roster_legend_docx(state: dict[str, Any], block: dict[str, Any]) -> bytes:
    document = Document()
    _title(document, block, "Roster and Legend")

    document.add_heading("Legend", level=2)
    legend = build_legend(state, block)
    legend_table = document.add_table(rows=1, cols=4)
    legend_table.style = "Table Grid"
    _set_cells(legend_table.rows[0].cells, ["#", "Name", "Active Dates", "Continuity Clinic"])
    for entry in legend.get("entries") or []:
        row = legend_table.add_row()
        _set_cells(
            row.cells,
            [
                entry.get("number"),
                entry.get("displayLabel"),
                entry.get("dateRange"),
                entry.get("continuityClinic"),
            ],
        )

    document.add_heading("Roster", level=2)
    roster_table = document.add_table(rows=1, cols=6)
    roster_table.style = "Table Grid"
    _set_cells(roster_table.rows[0].cells, ["Name", "Program", "Level", "Role", "Continuity", "Segments"])
    for rotator in state.get("rotators") or []:
        row = roster_table.add_row()
        _set_cells(
            row.cells,
            [
                rotator.get("displayName") or rotator.get("fullName"),
                rotator.get("program"),
                rotator.get("level"),
                rotator.get("role"),
                rotator.get("continuityClinic"),
                "\n".join(_segment_label(segment) for segment in rotator.get("segments") or []),
            ],
        )
    return _document_bytes(document)


def _title(document: Document, block: dict[str, Any], label: str) -> None:
    document.add_heading(f"{block.get('name') or 'Schedule'} - {label}", level=1)
    document.add_paragraph(f"{block.get('startDate') or ''} to {block.get('endDate') or ''}")


def _set_cells(cells: Any, values: list[Any]) -> None:
    for cell, value in zip(cells, values):
        cell.text = str(value or "")


def _document_bytes(document: Document) -> bytes:
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def _block_dates(block: dict[str, Any]) -> list[str]:
    if not block.get("startDate") or not block.get("endDate"):
        return []
    return date_range(block["startDate"], block["endDate"])


def _rotator_label(state: dict[str, Any], rotator_id: str | None) -> str:
    rotator = get_rotator(state, rotator_id)
    return (rotator or {}).get("displayName") or (rotator or {}).get("fullName") or rotator_id or "[Removed]"


def _clinic_line(row: dict[str, Any]) -> str:
    destination = row.get("clinicName") or "Clinic"
    attending = row.get("attendingName")
    if attending:
        destination = f"{destination} with {attending}"
    if row.get("stale"):
        destination = "Removed clinic"
    return f"{row.get('session') or ''} {destination}: {row.get('rotatorName') or ''}".strip()


def _segment_label(segment: dict[str, Any]) -> str:
    label = f"{segment.get('start') or ''} to {segment.get('end') or ''}"
    if segment.get("defaultPhase"):
        label += f" ({segment['defaultPhase']})"
    return label


def _slug(value: str) -> str:
    out = "".join(ch if ch.isalnum() else "-" for ch in value.strip().lower()).strip("-")
    while "--" in out:
        out = out.replace("--", "-")
    return out or "schedule"
