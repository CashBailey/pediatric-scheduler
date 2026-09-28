from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from backend_py.coordinator_docx_import import (
    NameResolver,
    default_coordinator_docx_paths,
    parse_coordinator_docx_bundle,
    parse_master_schedule,
)


def _shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def _write_master(path: Path, month: str = "July") -> Path:
    document = Document()
    document.add_paragraph(f"PediScheduler - {month} 2026 Master Schedule")
    table = document.add_table(rows=7, cols=4)
    headers = ["Rotator", "We\n1", "Th\n2", "Fr\n3"]
    for index, value in enumerate(headers):
        table.rows[0].cells[index].text = value

    rows = [
        ["\u2605 Jamie Rowe", "IP\nAM", "IP", "OP"],
        ["Blair Hart", "OP\nPM", "OFF", "OFF"],
        ["Parker Ross, MD", "", "IP\nPM", "IP"],
        ["Quinn Hayes, MD", "IP\nPM", "IP", "OFF"],
        ["Avery Stone", "", "OP", "OP"],
        ["Inpatient coverage", "2", "3", "1"],
    ]
    for row_index, values in enumerate(rows, start=1):
        for column_index, value in enumerate(values):
            table.rows[row_index].cells[column_index].text = value
    _shade(table.rows[5].cells[1], "bdbdbd")
    document.save(path)
    return path


def _write_inpatient(path: Path, month: str = "July") -> Path:
    document = Document()
    document.add_paragraph("\u2605 = Pediatric Neurology Fellow on service")
    table = document.add_table(rows=2, cols=7)
    for index, weekday in enumerate(["Sun", "Mon", "Tue", "Wed", "Thu", "Fri", "Sat"]):
        table.rows[0].cells[index].text = weekday
    table.rows[1].cells[3].text = (
        f"{month} 1\n"
        "\u2605 Coordinator AM Clinic\n"
        "Quinn Hayes PM AHD\n"
        "\u2605 FELLOW OFF"
    )
    table.rows[1].cells[4].text = f"{month} 2\nAngela Ross"
    document.save(path)
    return path


def _write_outpatient(path: Path, month: str = "July") -> Path:
    document = Document()
    document.add_paragraph(f"Week of {month} 1-3")
    table = document.add_table(rows=3, cols=4)
    headers = ["", f"Wednesday   \u00b7   {month} 1", f"Thursday   \u00b7   {month} 2", f"Friday   \u00b7   {month} 3"]
    for index, value in enumerate(headers):
        table.rows[0].cells[index].text = value
    table.rows[1].cells[0].text = "AM"
    table.rows[1].cells[1].text = "Elm: Blair"
    table.rows[1].cells[2].text = "GNU: \u2014"
    table.rows[1].cells[3].text = "Elm: Avery"
    table.rows[2].cells[0].text = "PM"
    table.rows[2].cells[1].text = "Continuity: Blair"
    table.rows[2].cells[2].text = "No clinic"
    table.rows[2].cells[3].text = "Next block \u2014 TBD"
    document.save(path)
    return path


def test_master_parser_preserves_periods_and_unavailable_fill(tmp_path):
    master = parse_master_schedule(_write_master(tmp_path / "master.docx"))

    coordinator = next(rotator for rotator in master["rotators"] if rotator["fullName"] == "Jamie Rowe")
    avery = next(rotator for rotator in master["rotators"] if rotator["fullName"] == "Avery Stone")

    assert master["counts"]["rotators"] == 5
    assert coordinator["isFellow"] is True
    assert coordinator["daily"][0]["status"] == "IP"
    assert coordinator["daily"][0]["period"] == "AM"
    assert avery["daily"][0]["status"] == "UNAVAILABLE"
    assert avery["unavailableRanges"] == [{"start": "2026-07-01", "end": "2026-07-01"}]
    assert master["coverageRows"]["Inpatient coverage"]["2026-07-02"] == 3


def test_bundle_probe_resolves_aliases_and_marks_half_day_mappings(tmp_path):
    preview = parse_coordinator_docx_bundle(
        master_path=_write_master(tmp_path / "master.docx"),
        inpatient_path=_write_inpatient(tmp_path / "inpatient.docx"),
        outpatient_path=_write_outpatient(tmp_path / "outpatient.docx"),
    )

    assert preview["acceptance"]["unresolvedNames"] == 0
    assert preview["acceptance"]["rotators"] == 5
    assert preview["acceptance"]["inpatientAssignments"] == 3
    assert preview["acceptance"]["outpatientSessions"] == 3
    assert preview["reconciliation"]["counts"]["mismatches"] == 0
    assert preview["inpatient"]["counts"]["halfDayAnnotations"] == 2
    assert preview["acceptance"]["halfDayFacts"] == 6

    facts = preview["schedulerStatePreview"]["halfDayFacts"]
    assert len(facts) == 6
    assert {fact["kind"] for fact in facts} == {"master-service-status", "inpatient-annotation"}
    assert any(
        fact["kind"] == "master-service-status"
        and fact["rotatorId"] == "rotator-jamie-rowe"
        and fact["date"] == "2026-07-01"
        and fact["period"] == "AM"
        and fact["status"] == "IP"
        for fact in facts
    )
    assert any(
        fact["kind"] == "inpatient-annotation"
        and fact["rotatorId"] == "rotator-quinn-hayes-md"
        and fact["date"] == "2026-07-01"
        and fact["period"] == "PM"
        and fact["label"] == "PM AHD"
        for fact in facts
    )
    assert any(item["reason"] == "FELLOW OFF" for item in preview["inpatient"]["skipped"])
    assert any(item["reason"] == "empty-slot" for item in preview["outpatient"]["skipped"])
    # DOCX bundles carry no program info — the import must say so up front.
    assert any(
        "5 rotators imported with program 'Other'" in warning for warning in preview["warnings"]
    )

    unsafe = [
        item
        for item in preview["schedulerStatePreview"]["inpatientAssignments"]
        if item.get("unsafeWholeDayMapping")
    ]
    assert {item["rotatorId"] for item in unsafe} == {
        "rotator-jamie-rowe",
        "rotator-quinn-hayes-md",
    }
    continuity = [
        item
        for item in preview["schedulerStatePreview"]["outpatientSessions"]
        if item["clinic"] == "Continuity"
    ]
    assert continuity[0]["provider"] == ""
    assert continuity[0]["rotatorId"] == "rotator-blair-hart"


def test_default_coordinator_docx_paths_discovers_latest_complete_downloads_bundle(tmp_path):
    downloads = tmp_path / "Downloads"
    downloads.mkdir()
    # Incomplete July bundle should be ignored in favor of complete August.
    (downloads / "July_2026_Master_Schedule.docx").write_bytes(b"placeholder")
    (downloads / "July_2026_Outpatient_Assignments.docx").write_bytes(b"placeholder")
    august_master = downloads / "August_2026_Master_Schedule.docx"
    august_inpatient = downloads / "August_2026_Inpatient_Schedule.docx"
    august_outpatient = downloads / "August_2026_Outpatient_Schedule.docx"
    for path in (august_master, august_inpatient, august_outpatient):
        path.write_bytes(b"placeholder")

    assert default_coordinator_docx_paths(downloads) == {
        "master": august_master,
        "inpatient": august_inpatient,
        "outpatient": august_outpatient,
    }


def test_bundle_probe_infers_august_dates_from_download_filenames(tmp_path):
    preview = parse_coordinator_docx_bundle(
        master_path=_write_master(tmp_path / "August_2026_Master_Schedule.docx", month="August"),
        inpatient_path=_write_inpatient(tmp_path / "August_2026_Inpatient_Schedule.docx", month="August"),
        outpatient_path=_write_outpatient(tmp_path / "August_2026_Outpatient_Schedule.docx", month="August"),
    )

    state = preview["schedulerStatePreview"]
    assert state["activeBlockId"] == "block-2026-08-coordinator-docx"
    assert state["serviceBlocks"][0]["name"] == "August 2026 Coordinator DOCX preview"
    assert state["serviceBlocks"][0]["startDate"] == "2026-08-01"
    assert preview["inpatient"]["records"][0]["date"] == "2026-08-01"
    assert preview["outpatient"]["sessions"][0]["date"] == "2026-08-01"


def test_name_resolver_handles_coordinator_roster_typos():
    resolver = NameResolver(
        [
            {"id": "rotator-jordan", "fullName": "Jordan Lee", "displayName": "Jordan Lee"},
            {"id": "rotator-parker", "fullName": "Parker Ross, MD", "displayName": "Parker Ross, MD"},
            {"id": "rotator-quinn", "fullName": "Quinn Hayes, MD", "displayName": "Quinn Hayes, MD"},
            {"id": "rotator-blair", "fullName": "Blair Hart", "displayName": "Blair Hart"},
        ]
    )

    assert resolver.resolve("Jordan Lee").rotator_id == "rotator-jordan"
    assert resolver.resolve("Parker Ross").rotator_id == "rotator-parker"
    assert resolver.resolve("Quinn Hayes").rotator_id == "rotator-quinn"
    assert resolver.resolve("Blair H.").rotator_id == "rotator-blair"
