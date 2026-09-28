"""POST /api/scheduler/command — the native client's single mutation route.

Covers the command envelope contract, state persistence after a mutating
command, and the HTTP error boundaries (400 bad JSON, 200-with-ok:false for
command-level failures).
"""

import base64
import io

from docx import Document as WordDocument
from fastapi.testclient import TestClient
import pytest

from backend_py.main import create_app


def _docx_text(encoded: str) -> str:
    data = base64.b64decode(encoded)
    document = WordDocument(io.BytesIO(data))
    values = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            values.extend(cell.text for cell in row.cells)
    return "\n".join(values)


def test_rules_patch_mutates_and_persists(tmp_path):
    with TestClient(create_app(data_dir=tmp_path)) as client:
        response = client.post(
            "/api/scheduler/command",
            json={"type": "rules.patch", "input": {"patch": {"maxConsecutiveInpatientDays": 9}}},
        )
        assert response.status_code == 200
        body = response.json()
        assert body["ok"] is True
        assert body["changed"] is True
        assert body["state"]["rules"]["maxConsecutiveInpatientDays"] == 9

        # Persisted: a fresh GET reflects the mutation.
        after = client.get("/api/scheduler/state").json()
        assert after["rules"]["maxConsecutiveInpatientDays"] == 9


def test_rules_patch_accepts_flat_input_shape(tmp_path):
    # inputOf merges sibling keys, so {type, patch} works like {type, input:{patch}}.
    with TestClient(create_app(data_dir=tmp_path)) as client:
        body = client.post(
            "/api/scheduler/command",
            json={"type": "rules.patch", "patch": {"honorNoClinicHolidays": False}},
        ).json()
        assert body["ok"] is True
        assert body["state"]["rules"]["honorNoClinicHolidays"] is False


def test_poster_settings_patch_mutates_and_persists(tmp_path):
    with TestClient(create_app(data_dir=tmp_path)) as client:
        body = client.post(
            "/api/scheduler/command",
            json={
                "type": "posterSettings.patch",
                "input": {
                    "patch": {
                        "programName": "Child Neuro Rotation",
                        "chief": "Dr. Lane",
                        "notes": ["Bring badge.", "Check clinic location."],
                        "locations": [
                            {"name": "North Clinic", "address": "123 Lane"},
                            {"name": "South Clinic", "address": "456 Road"},
                        ],
                        "tagline": "Thanks for caring for patients.",
                    }
                },
            },
        ).json()
        assert body["ok"] is True
        assert body["changed"] is True
        assert body["message"] == "Updated poster settings."
        assert body["data"]["patch"]["programName"] == "Child Neuro Rotation"
        settings = body["state"]["posterSettings"]
        assert settings["programName"] == "Child Neuro Rotation"
        assert settings["chief"] == "Dr. Lane"
        assert settings["notes"] == ["Bring badge.", "Check clinic location."]
        assert settings["locations"] == [
            {"name": "North Clinic", "address": "123 Lane"},
            {"name": "South Clinic", "address": "456 Road"},
        ]

        after = client.get("/api/scheduler/state").json()
        assert after["posterSettings"]["programName"] == "Child Neuro Rotation"
        assert after["posterSettings"]["chief"] == "Dr. Lane"
        assert after["posterSettings"]["notes"] == ["Bring badge.", "Check clinic location."]
        assert after["posterSettings"]["locations"] == [
            {"name": "North Clinic", "address": "123 Lane"},
            {"name": "South Clinic", "address": "456 Road"},
        ]
        assert after["posterSettings"]["tagline"] == "Thanks for caring for patients."


@pytest.mark.parametrize(
    ("command", "expected_field"),
    [
        ({"type": "rules.patch", "input": {"patch": "bad"}}, "patch"),
        ({"type": "posterSettings.patch", "input": {"patch": []}}, "patch"),
        ({"type": "posterSettings.patch", "input": {"patch": {"notes": "Bring badge"}}}, "notes"),
        (
            {
                "type": "posterSettings.patch",
                "input": {"patch": {"locations": [{"name": "North", "address": 42}]}},
            },
            "locations",
        ),
        ({"type": "attending.update", "input": {"name": "Dr. Alder", "patch": None}}, "patch"),
        ({"type": "block.update", "input": {"blockRef": "July Block", "patch": []}}, "patch"),
        ({"type": "rotator.update", "input": {"rotatorRef": "No One", "patch": "bad"}}, "patch"),
        ({"type": "source.add", "input": {"fileName": "Bad warnings", "importWarnings": [123]}}, "importWarnings"),
        ({"type": "source.add", "input": {"fileName": "Bad count", "importWarningCount": -1}}, "importWarningCount"),
    ],
)
def test_malformed_patch_and_source_payloads_return_command_errors(tmp_path, command, expected_field):
    with TestClient(create_app(data_dir=tmp_path)) as client:
        body = client.post("/api/scheduler/command", json=command).json()
        assert body["ok"] is False
        assert body["changed"] is False
        assert body["error"]["code"] == "invalid_type"
        assert body["error"]["field"] == expected_field


def test_attending_and_expected_source_commands_mutate_and_persist(tmp_path):
    with TestClient(create_app(data_dir=tmp_path)) as client:
        seeded = client.get("/api/scheduler/state").json()
        seeded["attendings"] = []
        seeded["expectedSourcePrograms"] = []
        assert client.post("/api/scheduler/state", json=seeded).json()["ok"] is True

        added = client.post(
            "/api/scheduler/command",
            json={"type": "attending.add", "input": {"name": "Dr. Alder"}},
        ).json()
        assert added["ok"] is True
        assert added["state"]["attendings"] == [{"name": "Dr. Alder", "recurringClinics": [], "oneOffDates": []}]

        duplicate = client.post(
            "/api/scheduler/command",
            json={"type": "attending.add", "input": {"name": "dr. alder"}},
        ).json()
        assert duplicate["ok"] is False
        assert duplicate["error"]["code"] == "no_effect"

        profiled = client.post(
            "/api/scheduler/command",
            json={
                "type": "attending.update",
                "input": {
                    "name": "Dr. Alder",
                    "patch": {
                        "recurringClinics": [
                            {
                                "id": "alder-mon-am",
                                "weekday": "Monday",
                                "period": "AM",
                                "clinicName": "Continuity",
                                "location": "South",
                                "capacity": 2,
                                "allowedRoles": ["Resident", "Fellow"],
                            }
                        ],
                        "oneOffDates": [
                            {
                                "id": "alder-special",
                                "date": "2026-05-13",
                                "period": "PM",
                                "clinicName": "Makeup",
                                "location": "North",
                                "capacity": "1",
                                "allowedRoles": ["Student"],
                            }
                        ],
                    },
                },
            },
        ).json()
        assert profiled["ok"] is True
        assert profiled["state"]["attendings"][0]["recurringClinics"] == [
            {
                "id": "alder-mon-am",
                "weekday": "Monday",
                "period": "AM",
                "clinicName": "Continuity",
                "location": "South",
                "capacity": 2,
                "allowedRoles": ["Resident", "Fellow"],
            }
        ]
        assert profiled["state"]["attendings"][0]["oneOffDates"] == [
            {
                "id": "alder-special",
                "date": "2026-05-13",
                "period": "PM",
                "clinicName": "Makeup",
                "location": "North",
                "capacity": "1",
                "allowedRoles": ["Student"],
            }
        ]

        updated = client.post(
            "/api/scheduler/command",
            json={"type": "attending.update", "input": {"name": "Dr. Alder", "patch": {"name": "Dr. Cedar"}}},
        ).json()
        assert updated["ok"] is True
        assert updated["state"]["attendings"][0]["name"] == "Dr. Cedar"

        source = client.post(
            "/api/scheduler/command",
            json={"type": "expectedSource.add", "input": {"program": "Methodist"}},
        ).json()
        assert source["ok"] is True
        assert source["state"]["expectedSourcePrograms"] == ["Methodist"]

        source_again = client.post(
            "/api/scheduler/command",
            json={"type": "expectedSource.add", "input": {"program": "Methodist"}},
        ).json()
        assert source_again["ok"] is True
        assert source_again["state"]["expectedSourcePrograms"] == ["Methodist"]

        removed = client.post(
            "/api/scheduler/command",
            json={"type": "attending.remove", "input": {"name": "Dr. Cedar"}},
        ).json()
        assert removed["ok"] is True
        assert removed["state"]["attendings"] == []

        source_removed = client.post(
            "/api/scheduler/command",
            json={"type": "expectedSource.remove", "input": {"program": "Methodist"}},
        ).json()
        assert source_removed["ok"] is True
        assert source_removed["state"]["expectedSourcePrograms"] == []

        after = client.get("/api/scheduler/state").json()
        assert after["attendings"] == []
        assert after["expectedSourcePrograms"] == []


@pytest.mark.parametrize("command_type", ["source.delete", "source.remove"])
def test_source_delete_removes_record_without_cascading_imported_people(tmp_path, command_type):
    with TestClient(create_app(data_dir=tmp_path)) as client:
        seeded = client.get("/api/scheduler/state").json()
        seeded.update(
            {
                "sources": [
                    {
                        "id": "source-july-roster",
                        "importedAt": "2026-07-05",
                        "status": "Reviewed",
                        "program": "UT Pediatrics",
                        "fileName": "July roster.csv",
                        "fileType": "csv",
                        "content": "",
                        "parsedRows": [{"Name": "Drew Quinn", "Program": "UT Pediatrics"}],
                        "importedRotatorCount": 1,
                    }
                ],
                "rotators": [
                    {
                        "id": "rot-ari",
                        "fullName": "Drew Quinn",
                        "displayName": "Drew Quinn",
                        "program": "UT Pediatrics",
                        "level": "PGY-2",
                        "role": "Resident",
                        "segments": [{"start": "2026-07-01", "end": "2026-07-31"}],
                        "schoolType": "ut-peds",
                        "continuityClinic": "",
                        "dayOff": [],
                        "unavailableRanges": [],
                    }
                ],
                "inpatientAssignments": [
                    {"id": "ip-ari", "date": "2026-07-06", "rotatorId": "rot-ari", "role": "Resident", "source": "Manual"}
                ],
                "outpatientSessions": [
                    {
                        "id": "op-ari",
                        "date": "2026-07-07",
                        "period": "AM",
                        "clinic": "Continuity",
                        "rotatorId": "rot-ari",
                    }
                ],
            }
        )
        assert client.post("/api/scheduler/state", json=seeded).json()["ok"] is True

        body = client.post(
            "/api/scheduler/command",
            json={"type": command_type, "input": {"sourceId": "source-july-roster"}},
        ).json()
        assert body["ok"] is True
        assert body["changed"] is True
        assert body["state"]["sources"] == []
        assert body["state"]["rotators"] == seeded["rotators"]
        assert body["state"]["inpatientAssignments"] == seeded["inpatientAssignments"]
        assert body["state"]["outpatientSessions"] == seeded["outpatientSessions"]

        after = client.get("/api/scheduler/state").json()
        assert after["sources"] == []
        assert after["rotators"][0]["fullName"] == "Drew Quinn"


def test_source_delete_rejects_unknown_source_without_mutating(tmp_path):
    with TestClient(create_app(data_dir=tmp_path)) as client:
        seeded = client.get("/api/scheduler/state").json()
        seeded["sources"] = [{"id": "source-keep", "fileName": "keep.csv"}]
        assert client.post("/api/scheduler/state", json=seeded).json()["ok"] is True

        body = client.post(
            "/api/scheduler/command",
            json={"type": "source.delete", "input": {"sourceId": "missing"}},
        ).json()
        assert body["ok"] is False
        assert body["changed"] is False
        assert body["error"]["code"] == "source_not_found"

        after = client.get("/api/scheduler/state").json()
        assert after["sources"] == [{"id": "source-keep", "fileName": "keep.csv"}]


def test_source_add_manual_note_persists_reviewed_source(tmp_path):
    with TestClient(create_app(data_dir=tmp_path)) as client:
        response = client.post(
            "/api/scheduler/command",
            json={
                "type": "source.add",
                "input": {
                    "id": "source-manual-coordinator-note",
                    "program": "Other",
                    "fileName": "Coordinator note",
                    "content": "Use the updated fellow preference sheet.",
                },
            },
        )
        assert response.status_code == 200
        body = response.json()

        assert body["ok"] is True
        assert body["changed"] is True
        assert body["data"]["sourceId"] == "source-manual-coordinator-note"
        source = body["state"]["sources"][0]
        assert source["id"] == "source-manual-coordinator-note"
        assert source["fileName"] == "Coordinator note"
        assert source["fileType"] == "manual"
        assert source["status"] == "Reviewed"
        assert source["program"] == "Other"
        assert source["content"] == "Use the updated fellow preference sheet."
        assert source["parsedRows"] == []

        after = client.get("/api/scheduler/state").json()
        assert after["sources"] == body["state"]["sources"]

        duplicate = client.post(
            "/api/scheduler/command",
            json={"type": "source.add", "input": {"id": "source-manual-coordinator-note", "fileName": "Duplicate"}},
        ).json()
        assert duplicate["ok"] is False
        assert duplicate["error"]["code"] == "duplicate_source"


def test_block_commands_add_use_update_delete_and_persist(tmp_path):
    with TestClient(create_app(data_dir=tmp_path)) as client:
        initial = client.get("/api/scheduler/state").json()
        original_id = initial["activeBlockId"]

        added = client.post(
            "/api/scheduler/command",
            json={
                "type": "block.add",
                "input": {
                    "name": "June Block",
                    "startDate": "2026-06-01",
                    "endDate": "2026-06-28",
                },
            },
        ).json()
        assert added["ok"] is True
        assert added["changed"] is True
        new_id = added["data"]["blockId"]
        assert new_id.startswith("block-")
        assert added["state"]["activeBlockId"] == new_id
        assert added["state"]["serviceBlocks"][-1]["name"] == "June Block"
        assert added["state"]["serviceBlocks"][-1]["generate"]["export"] is True

        use_original = client.post(
            "/api/scheduler/command",
            json={"type": "block.use", "input": {"blockRef": original_id}},
        ).json()
        assert use_original["ok"] is True
        assert use_original["state"]["activeBlockId"] == original_id

        updated = client.post(
            "/api/scheduler/command",
            json={
                "type": "block.update",
                "input": {
                    "blockRef": new_id,
                    "patch": {
                        "name": "Renamed June",
                        "startDate": "2026-06-02",
                        "endDate": "2026-06-29",
                        "coverage": {
                            "weekday": {"ip": {"count": 3}},
                            "saturday": {"ip": {"count": 1}},
                            "sunday": {"ip": {"count": 1}},
                            "holiday": {"ip": {"count": 2}},
                        },
                        "holidays": [{"date": "2026-06-19", "label": "Juneteenth", "noClinic": True}],
                    },
                },
            },
        ).json()
        assert updated["ok"] is True
        assert updated["state"]["activeBlockId"] == original_id
        renamed = next(block for block in updated["state"]["serviceBlocks"] if block["id"] == new_id)
        assert renamed["name"] == "Renamed June"
        assert renamed["startDate"] == "2026-06-02"
        assert renamed["endDate"] == "2026-06-29"
        assert renamed["coverage"]["weekday"]["ip"]["count"] == 3
        assert renamed["holidays"] == [{"date": "2026-06-19", "label": "Juneteenth", "noClinic": True}]

        finalized = client.post(
            "/api/scheduler/command",
            json={
                "type": "block.update",
                "input": {
                    "blockRef": new_id,
                    "patch": {
                        "status": "Final",
                        "finalizedAt": "2026-07-06T12:00:00.000Z",
                        "finalizedBy": "Native Reports Review & Finalize",
                        "finalReview": {
                            "reviewedAt": "2026-07-06T12:00:00.000Z",
                            "criticalConflictCount": 0,
                            "warningConflictCount": 2,
                            "openSlots": 0,
                            "missingSourcePrograms": [],
                            "reason": "Reports Review & Finalize checks passed",
                        },
                    },
                },
            },
        ).json()
        assert finalized["ok"] is True
        final_block = next(block for block in finalized["state"]["serviceBlocks"] if block["id"] == new_id)
        assert final_block["status"] == "Final"
        assert final_block["finalizedAt"] == "2026-07-06T12:00:00.000Z"
        assert final_block["finalizedBy"] == "Native Reports Review & Finalize"
        assert final_block["finalReview"]["criticalConflictCount"] == 0
        assert final_block["finalReview"]["warningConflictCount"] == 2
        assert final_block["finalReview"]["missingSourcePrograms"] == []
        assert "postFinalChanges" not in final_block

        post_final = client.post(
            "/api/scheduler/command",
            json={
                "type": "block.update",
                "input": {
                    "blockRef": new_id,
                    "patch": {"name": "Renamed June Final"},
                    "postFinalReason": "Corrected display title after final review",
                    "changedBy": "Command test",
                },
            },
        ).json()
        assert post_final["ok"] is True
        post_final_block = next(block for block in post_final["state"]["serviceBlocks"] if block["id"] == new_id)
        assert post_final_block["name"] == "Renamed June Final"
        post_final_changes = post_final_block["postFinalChanges"]
        assert len(post_final_changes) == 1
        assert post_final_changes[0]["blockId"] == new_id
        assert post_final_changes[0]["command"] == "block.update"
        assert post_final_changes[0]["reason"] == "Corrected display title after final review"
        assert post_final_changes[0]["changedBy"] == "Command test"
        assert post_final_changes[0]["summary"] == "Updated block Renamed June Final."
        assert post_final["data"]["postFinalChange"] == post_final_changes[0]

        persisted = client.get("/api/scheduler/state").json()
        persisted_block = next(block for block in persisted["serviceBlocks"] if block["id"] == new_id)
        assert persisted_block["postFinalChanges"] == post_final_changes

        use_renamed = client.post(
            "/api/scheduler/command",
            json={"type": "block.use", "input": {"name": "Renamed June Final"}},
        ).json()
        assert use_renamed["ok"] is True
        assert use_renamed["state"]["activeBlockId"] == new_id

        deleted = client.post(
            "/api/scheduler/command",
            json={"type": "block.delete", "input": {"blockRef": new_id}},
        ).json()
        assert deleted["ok"] is True
        assert deleted["state"]["activeBlockId"] == original_id
        assert [block["id"] for block in deleted["state"]["serviceBlocks"]] == [original_id]

        after = client.get("/api/scheduler/state").json()
        assert after["activeBlockId"] == original_id
        assert [block["id"] for block in after["serviceBlocks"]] == [original_id]


def test_block_commands_validate_inputs(tmp_path):
    with TestClient(create_app(data_dir=tmp_path)) as client:
        partial_add = client.post(
            "/api/scheduler/command",
            json={"type": "block.add", "input": {"name": "Partial"}},
        ).json()
        assert partial_add["ok"] is False
        assert partial_add["error"]["code"] == "missing_field"
        assert partial_add["error"]["field"] == "startDate"

        bad_update = client.post(
            "/api/scheduler/command",
            json={
                "type": "block.update",
                "input": {"blockRef": "block-new", "patch": {"startDate": "June"}},
            },
        ).json()
        assert bad_update["ok"] is False
        assert bad_update["error"]["code"] == "invalid_date"
        assert bad_update["error"]["field"] == "startDate"

        only_block_delete = client.post(
            "/api/scheduler/command",
            json={"type": "block.delete", "input": {"blockRef": "block-new"}},
        ).json()
        assert only_block_delete["ok"] is False
        assert only_block_delete["error"]["code"] == "no_effect"
        assert "only block" in only_block_delete["error"]["message"]


def test_grid_show_command_returns_planning_grid_without_mutation(tmp_path):
    with TestClient(create_app(data_dir=tmp_path)) as client:
        state = client.get("/api/scheduler/state").json()
        block = {
            "id": "b1",
            "name": "Grid Block",
            "startDate": "2026-05-04",
            "endDate": "2026-05-06",
            "status": "Draft",
            "generate": {},
            "holidays": [],
        }
        state.update(
            {
                "activeBlockId": "b1",
                "serviceBlocks": [block],
                "rotators": [
                    {
                        "id": "r1",
                        "fullName": "Drew Quinn",
                        "displayName": "Drew Quinn",
                        "program": "UT Pediatrics",
                        "level": "PGY-2",
                        "role": "Resident",
                        "segments": [{"start": "2026-05-04", "end": "2026-05-06"}],
                        "schoolType": "ut-peds",
                        "continuityClinic": "",
                        "dayOff": [],
                        "unavailableRanges": [],
                    },
                    {
                        "id": "r2",
                        "fullName": "Maya Lopez",
                        "displayName": "Maya Lopez",
                        "program": "UT Pediatrics",
                        "level": "PGY-3",
                        "role": "Resident",
                        "segments": [{"start": "2026-05-04", "end": "2026-05-06"}],
                        "schoolType": "ut-peds",
                        "continuityClinic": "",
                        "dayOff": ["Tuesday"],
                        "unavailableRanges": [],
                    },
                    {
                        "id": "r3",
                        "fullName": "Noah Park",
                        "displayName": "Noah Park",
                        "program": "UT Pediatrics",
                        "level": "PGY-4",
                        "role": "Resident",
                        "segments": [],
                        "schoolType": "ut-peds",
                        "continuityClinic": "",
                        "dayOff": [],
                        "unavailableRanges": [],
                    },
                ],
                "inpatientAssignments": [
                    {"id": "ip-r1-mon", "date": "2026-05-04", "rotatorId": "r1", "role": "Resident"},
                    {"id": "off-r2-mon", "date": "2026-05-04", "rotatorId": "r2", "role": "Off"},
                    {"id": "ip-r1-wed", "date": "2026-05-06", "rotatorId": "r1", "role": "Resident"},
                ],
                "outpatientSessions": [
                    {"id": "op-r1-tue", "date": "2026-05-05", "period": "AM", "clinic": "QRS", "rotatorId": "r1"},
                    {"id": "op-r1-wed", "date": "2026-05-06", "period": "PM", "clinic": "TSC", "rotatorId": "r1"},
                ],
            }
        )
        assert client.post("/api/scheduler/state", json=state).json()["ok"] is True

        body = client.post("/api/scheduler/command", json={"type": "grid.show"}).json()
        assert body["ok"] is True
        assert body["changed"] is False
        assert body["message"] == "Built planning grid for Grid Block."
        assert body["data"]["blockId"] == "b1"

        grid = body["data"]["grid"]
        assert grid["dates"] == ["2026-05-04", "2026-05-05", "2026-05-06"]
        rows = {row["rotator"]["id"]: row for row in grid["rows"]}
        assert [cell["status"] for cell in rows["r1"]["cells"]] == ["inpatient", "outpatient", "both"]
        assert [cell["status"] for cell in rows["r2"]["cells"]] == ["off", "off", "unassigned"]
        assert [cell["status"] for cell in rows["r3"]["cells"]] == ["absent", "absent", "absent"]
        assert grid["totals"] == [
            {"date": "2026-05-04", "ip": 1, "op": 0, "both": 0, "unassigned": 0, "present": 1},
            {"date": "2026-05-05", "ip": 0, "op": 1, "both": 0, "unassigned": 0, "present": 1},
            {"date": "2026-05-06", "ip": 1, "op": 1, "both": 1, "unassigned": 1, "present": 2},
        ]
        assert [row["rotator"]["id"] for row in grid["sections"]["needs"]] == ["r2"]
        assert [row["rotator"]["id"] for row in grid["sections"]["mixed"]] == ["r1"]
        assert [row["rotator"]["id"] for row in grid["sections"]["unavailable"]] == ["r3"]

        after = client.get("/api/scheduler/state").json()
        assert after["inpatientAssignments"] == state["inpatientAssignments"]
        assert after["outpatientSessions"] == state["outpatientSessions"]


def test_grid_show_peek_extends_projection_without_mutation(tmp_path):
    with TestClient(create_app(data_dir=tmp_path)) as client:
        state = client.get("/api/scheduler/state").json()
        state.update(
            {
                "activeBlockId": "b1",
                "serviceBlocks": [
                    {
                        "id": "b0",
                        "name": "June",
                        "startDate": "2026-06-01",
                        "endDate": "2026-06-30",
                        "status": "Draft",
                        "generate": {},
                        "holidays": [],
                    },
                    {
                        "id": "b1",
                        "name": "July",
                        "startDate": "2026-07-01",
                        "endDate": "2026-07-31",
                        "status": "Draft",
                        "generate": {},
                        "holidays": [],
                    },
                    {
                        "id": "b2",
                        "name": "August",
                        "startDate": "2026-08-01",
                        "endDate": "2026-08-31",
                        "status": "Draft",
                        "generate": {},
                        "holidays": [],
                    },
                ],
                "rotators": [
                    {
                        "id": "r1",
                        "fullName": "Maya Lopez",
                        "displayName": "Maya Lopez",
                        "program": "Methodist",
                        "level": "PGY-3",
                        "role": "Resident",
                        "segments": [{"start": "2026-07-01", "end": "2026-07-31"}],
                        "schoolType": "methodist",
                        "continuityClinic": "",
                        "dayOff": [],
                        "unavailableRanges": [],
                    }
                ],
                "inpatientAssignments": [],
                "outpatientSessions": [],
            }
        )
        assert client.post("/api/scheduler/state", json=state).json()["ok"] is True

        off = client.post("/api/scheduler/command", json={"type": "grid.show"}).json()
        assert off["ok"] is True
        assert off["changed"] is False
        assert off["data"]["peekBeforeBlock"] is False
        assert off["data"]["peekPastBlock"] is False
        assert off["data"]["rawStartDate"] == "2026-07-01"
        assert off["data"]["rawEndDate"] == "2026-07-31"
        assert off["data"]["effectiveStartDate"] == "2026-07-01"
        assert off["data"]["effectiveEndDate"] == "2026-07-31"
        assert off["data"]["grid"]["dates"][-1] == "2026-07-31"

        on = client.post(
            "/api/scheduler/command",
            json={"type": "grid.show", "input": {"peekPastBlock": True}},
        ).json()
        assert on["ok"] is True
        assert on["changed"] is False
        assert on["data"]["peekPastBlock"] is True
        assert on["data"]["rawEndDate"] == "2026-07-31"
        assert on["data"]["effectiveEndDate"] == "2026-08-14"
        assert on["data"]["grid"]["dates"][0] == "2026-07-01"
        assert on["data"]["grid"]["dates"][-1] == "2026-08-14"

        prior = client.post(
            "/api/scheduler/command",
            json={"type": "grid.show", "input": {"peekBeforeBlock": True}},
        ).json()
        assert prior["ok"] is True
        assert prior["data"]["peekBeforeBlock"] is True
        assert prior["data"]["rawStartDate"] == "2026-07-01"
        assert prior["data"]["effectiveStartDate"] == "2026-06-17"
        assert prior["data"]["grid"]["dates"][0] == "2026-06-17"

        after = client.get("/api/scheduler/state").json()
        assert after["serviceBlocks"] == state["serviceBlocks"]
        assert after["inpatientAssignments"] == []
        assert after["outpatientSessions"] == []


def test_grid_show_peek_uses_fixed_fourteen_day_buffer_despite_overlapping_block(tmp_path):
    with TestClient(create_app(data_dir=tmp_path)) as client:
        state = client.get("/api/scheduler/state").json()
        state.update(
            {
                "activeBlockId": "b1",
                "serviceBlocks": [
                    {
                        "id": "b1",
                        "name": "July",
                        "startDate": "2026-07-01",
                        "endDate": "2026-07-31",
                        "status": "Draft",
                        "generate": {},
                        "holidays": [],
                    },
                    {
                        "id": "b2",
                        "name": "Overlap",
                        "startDate": "2026-07-15",
                        "endDate": "2026-07-31",
                        "status": "Draft",
                        "generate": {},
                        "holidays": [],
                    },
                ],
                "rotators": [
                    {
                        "id": "r1",
                        "fullName": "Maya Lopez",
                        "displayName": "Maya Lopez",
                        "program": "Methodist",
                        "level": "PGY-3",
                        "role": "Resident",
                        "segments": [{"start": "2026-07-01", "end": "2026-07-31"}],
                        "schoolType": "methodist",
                        "continuityClinic": "",
                        "dayOff": [],
                        "unavailableRanges": [],
                    }
                ],
                "inpatientAssignments": [],
                "outpatientSessions": [],
            }
        )
        assert client.post("/api/scheduler/state", json=state).json()["ok"] is True

        body = client.post(
            "/api/scheduler/command",
            json={"type": "grid.show", "input": {"peekPastBlock": True}},
        ).json()
        assert body["ok"] is True
        assert body["changed"] is False
        assert body["data"]["peekPastBlock"] is True
        assert body["data"]["rawEndDate"] == "2026-07-31"
        assert body["data"]["effectiveEndDate"] == "2026-08-14"
        assert body["data"]["grid"]["dates"][-14:] == [
            "2026-08-01",
            "2026-08-02",
            "2026-08-03",
            "2026-08-04",
            "2026-08-05",
            "2026-08-06",
            "2026-08-07",
            "2026-08-08",
            "2026-08-09",
            "2026-08-10",
            "2026-08-11",
            "2026-08-12",
            "2026-08-13",
            "2026-08-14",
        ]


@pytest.mark.parametrize("command_type", ["draft.generate", "inpatient.draft"])
def test_inpatient_draft_command_generates_and_persists(tmp_path, command_type):
    with TestClient(create_app(data_dir=tmp_path)) as client:
        state = client.get("/api/scheduler/state").json()
        block = {
            "id": "b1",
            "name": "Draftable",
            "startDate": "2026-05-04",
            "endDate": "2026-05-08",
            "status": "Draft",
            "generate": {},
            "holidays": [],
            "coverage": {"weekday": {"ip": {"count": 1}}},
        }
        state.update(
            {
                "activeBlockId": "b1",
                "serviceBlocks": [block],
                "rotators": [
                    {
                        "id": "r1",
                        "fullName": "Drew Quinn",
                        "displayName": "Drew Quinn",
                        "program": "Other",
                        "level": "PGY-2",
                        "role": "Resident",
                        "segments": [{"start": "2026-05-04", "end": "2026-05-08"}],
                        "schoolType": "other",
                        "continuityClinic": "",
                        "dayOff": [],
                        "unavailableRanges": [],
                    }
                ],
                "inpatientAssignments": [],
                "outpatientSessions": [],
            }
        )
        assert client.post("/api/scheduler/state", json=state).json()["ok"] is True

        response = client.post("/api/scheduler/command", json={"type": command_type})
        assert response.status_code == 200
        body = response.json()
        assert body["ok"] is True
        assert body["changed"] is True
        assert body["data"]["inpatientAdded"] == 5
        assert body["data"]["report"]["summary"]["inpatientAdded"] == 5
        assert len([item for item in body["state"]["inpatientAssignments"] if item["source"] == "Auto-Draft"]) == 5

        after = client.get("/api/scheduler/state").json()
        assert len(after["inpatientAssignments"]) == 5


def test_draft_command_applies_segment_preassignments(tmp_path):
    with TestClient(create_app(data_dir=tmp_path)) as client:
        state = client.get("/api/scheduler/state").json()
        block = {
            "id": "b1",
            "name": "Template Block",
            "startDate": "2026-09-14",
            "endDate": "2026-09-18",
            "status": "Draft",
            "generate": {},
            "holidays": [{"date": "2026-09-16", "noClinic": True, "label": "Closed clinic"}],
            "coverage": {
                "weekday": {"ip": {"count": 0}},
                "saturday": {"ip": {"count": 0}},
                "sunday": {"ip": {"count": 0}},
                "holiday": {"ip": {"count": 0}},
            },
        }
        state.update(
            {
                "activeBlockId": "b1",
                "serviceBlocks": [block],
                "rotators": [
                    {
                        "id": "r1",
                        "fullName": "Drew Quinn",
                        "displayName": "Drew Quinn",
                        "program": "Other",
                        "level": "PGY-2",
                        "role": "Resident",
                        "segments": [{"start": "2026-09-14", "end": "2026-09-18", "defaultPhase": "outpatient"}],
                        "schoolType": "other",
                        "continuityClinic": "Tuesday PM",
                        "dayOff": [],
                        "unavailableRanges": [],
                    },
                    {
                        "id": "r2",
                        "fullName": "Bea Lee",
                        "displayName": "Bea Lee",
                        "program": "Other",
                        "level": "PGY-3",
                        "role": "Resident",
                        "segments": [{"start": "2026-09-14", "end": "2026-09-18", "defaultPhase": "inpatient"}],
                        "schoolType": "other",
                        "continuityClinic": "",
                        "dayOff": [],
                        "unavailableRanges": [],
                    },
                ],
                "inpatientAssignments": [
                    {"id": "manual-ip", "date": "2026-09-15", "rotatorId": "r2", "role": "Resident", "source": "Manual"}
                ],
                "outpatientSessions": [],
            }
        )
        assert client.post("/api/scheduler/state", json=state).json()["ok"] is True

        response = client.post("/api/scheduler/command", json={"type": "draft.generate"})
        assert response.status_code == 200
        body = response.json()

        assert body["ok"] is True
        assert body["data"]["inpatientAdded"] == 4
        assert body["data"]["outpatientAdded"] == 7
        # The nested report must agree with the top-level counts: it used to
        # baseline against the already-preassigned state and claim 0/0.
        assert body["data"]["report"]["summary"]["inpatientAdded"] == 4
        assert body["data"]["report"]["summary"]["outpatientAdded"] == 7
        pre_ip = [item for item in body["state"]["inpatientAssignments"] if item["source"] == "Auto-Preassigned"]
        pre_op = [item for item in body["state"]["outpatientSessions"] if item["source"] == "Auto-Preassigned"]
        assert len(pre_ip) == 4
        assert len(pre_op) == 7
        assert {item["date"] for item in pre_ip} == {"2026-09-14", "2026-09-16", "2026-09-17", "2026-09-18"}
        assert not any(item["date"] == "2026-09-16" for item in pre_op)
        assert not any(item["date"] == "2026-09-15" and item["period"] == "PM" for item in pre_op)

        again = client.post("/api/scheduler/command", json={"type": "draft.generate"}).json()
        assert again["ok"] is True
        assert again["data"]["inpatientAdded"] == 0
        assert again["data"]["outpatientAdded"] == 0


def test_methodist_auto_command_generates_14_14_and_is_idempotent(tmp_path):
    with TestClient(create_app(data_dir=tmp_path)) as client:
        state = client.get("/api/scheduler/state").json()
        block = {
            "id": "b1",
            "name": "Methodist Block",
            "startDate": "2026-06-20",
            "endDate": "2026-07-17",
            "status": "Draft",
            "generate": {},
            "holidays": [],
            "coverage": {
                "weekday": {"ip": {"count": 0}},
                "saturday": {"ip": {"count": 0}},
                "sunday": {"ip": {"count": 0}},
                "holiday": {"ip": {"count": 0}},
            },
        }
        state.update(
            {
                "activeBlockId": "b1",
                "serviceBlocks": [block],
                "rotators": [
                    {
                        "id": "m1",
                        "fullName": "Maya Lopez",
                        "displayName": "Maya Lopez",
                        "program": "Methodist",
                        "level": "PGY-3",
                        "role": "Resident",
                        "segments": [{"start": "2026-06-20", "end": "2026-07-17"}],
                        "schoolType": "methodist",
                        "rotationStartDate": "2026-06-20",
                        "continuityClinic": "",
                        "dayOff": [],
                        "unavailableRanges": [],
                    }
                ],
                "inpatientAssignments": [],
                "outpatientSessions": [],
            }
        )
        assert client.post("/api/scheduler/state", json=state).json()["ok"] is True

        body = client.post("/api/scheduler/command", json={"type": "methodist.auto"}).json()
        assert body["ok"] is True
        assert body["changed"] is True
        assert body["data"]["methodistCount"] == 1
        assert body["data"]["startSidesSet"] == 1
        assert body["data"]["inpatientAdded"] > 0
        assert body["data"]["outpatientAdded"] > 0
        assert body["state"]["rotators"][0]["methodistStartSide"] == "outpatient"
        assert any(item["date"] == "2026-06-22" and item["rotatorId"] == "m1" for item in body["state"]["outpatientSessions"])
        assert any(item["date"] == "2026-07-06" and item["rotatorId"] == "m1" for item in body["state"]["inpatientAssignments"])

        again = client.post("/api/scheduler/command", json={"type": "methodist.auto"}).json()
        assert again["ok"] is True
        assert again["changed"] is False
        assert again["data"]["inpatientAdded"] == 0
        assert again["data"]["outpatientAdded"] == 0


def test_assign_range_command_generates_and_persists(tmp_path):
    with TestClient(create_app(data_dir=tmp_path)) as client:
        state = client.get("/api/scheduler/state").json()
        block = {
            "id": "b1",
            "name": "Editable",
            "startDate": "2026-05-04",
            "endDate": "2026-05-08",
            "status": "Draft",
            "generate": {},
            "holidays": [],
        }
        state.update(
            {
                "activeBlockId": "b1",
                "serviceBlocks": [block],
                "rotators": [
                    {
                        "id": "r1",
                        "fullName": "Drew Quinn",
                        "displayName": "Drew Quinn",
                        "program": "Other",
                        "level": "PGY-2",
                        "role": "Resident",
                        "segments": [{"start": "2026-05-04", "end": "2026-05-08"}],
                        "schoolType": "other",
                        "continuityClinic": "",
                        "dayOff": [],
                        "unavailableRanges": [],
                    }
                ],
                "inpatientAssignments": [],
                "outpatientSessions": [],
            }
        )
        assert client.post("/api/scheduler/state", json=state).json()["ok"] is True

        response = client.post(
            "/api/scheduler/command",
            json={
                "type": "assign.range",
                "input": {
                    "rotatorRef": "r1",
                    "startDate": "2026-05-04",
                    "endDate": "2026-05-05",
                    "phase": "outpatient",
                },
            },
        )
        assert response.status_code == 200
        body = response.json()
        assert body["ok"] is True
        assert body["changed"] is True
        assert body["data"]["rotatorId"] == "r1"
        assert body["data"]["phase"] == "outpatient"
        assert len([item for item in body["state"]["outpatientSessions"] if item["source"] == "Range-Assigned"]) == 4

        after = client.get("/api/scheduler/state").json()
        assert len(after["outpatientSessions"]) == 4


def test_assign_range_rejects_invalid_phase(tmp_path):
    with TestClient(create_app(data_dir=tmp_path)) as client:
        response = client.post(
            "/api/scheduler/command",
            json={
                "type": "assign.range",
                "input": {
                    "rotatorRef": "missing",
                    "startDate": "2026-05-04",
                    "endDate": "2026-05-05",
                    "phase": "vacation",
                },
            },
        )
        body = response.json()
        assert response.status_code == 200
        assert body["ok"] is False
        assert body["error"]["code"] == "invalid_enum"


def test_inpatient_assign_command_preserves_same_day_same_role_rotators(tmp_path):
    with TestClient(create_app(data_dir=tmp_path)) as client:
        state = client.get("/api/scheduler/state").json()
        state.update(
            {
                "activeBlockId": "b1",
                "serviceBlocks": [
                    {
                        "id": "b1",
                        "name": "May 2026",
                        "startDate": "2026-05-04",
                        "endDate": "2026-05-31",
                        "status": "Draft",
                        "generate": {},
                        "holidays": [],
                    }
                ],
                "rotators": [
                    {
                        "id": "r1",
                        "fullName": "Drew Quinn",
                        "displayName": "Drew Quinn",
                        "program": "Other",
                        "level": "PGY-2",
                        "role": "Resident",
                        "segments": [{"start": "2026-05-04", "end": "2026-05-08"}],
                        "schoolType": "other",
                        "continuityClinic": "",
                        "dayOff": [],
                        "unavailableRanges": [],
                    },
                    {
                        "id": "r2",
                        "fullName": "Blake Lee",
                        "displayName": "Blake Lee",
                        "program": "Other",
                        "level": "PGY-3",
                        "role": "Resident",
                        "segments": [{"start": "2026-05-04", "end": "2026-05-08"}],
                        "schoolType": "other",
                        "continuityClinic": "",
                        "dayOff": [],
                        "unavailableRanges": [],
                    },
                ],
                "inpatientAssignments": [],
            }
        )
        assert client.post("/api/scheduler/state", json=state).json()["ok"] is True

        first = client.post(
            "/api/scheduler/command",
            json={
                "type": "inpatient.assign",
                "input": {"rotatorRef": "Drew Quinn", "date": "2026-05-04", "source": "Drag-Drop"},
            },
        ).json()
        assert first["ok"] is True
        assert first["changed"] is True
        assert first["data"] == {"rotatorId": "r1", "date": "2026-05-04", "role": "Resident"}
        assert first["state"]["inpatientAssignments"] == [
            {
                "id": "in-2026-05-04-r1-resident",
                "date": "2026-05-04",
                "rotatorId": "r1",
                "role": "Resident",
                "source": "Drag-Drop",
            }
        ]

        second = client.post(
            "/api/scheduler/command",
            json={"type": "inpatient.assign", "input": {"rotatorRef": "r2", "date": "2026-05-04"}},
        ).json()
        assert second["ok"] is True
        assignments = second["state"]["inpatientAssignments"]
        assert len(assignments) == 2
        assert sorted(item["rotatorId"] for item in assignments) == ["r1", "r2"]
        assert next(item for item in assignments if item["rotatorId"] == "r2")["source"] == "Manual"

        again = client.post(
            "/api/scheduler/command",
            json={"type": "inpatient.assign", "input": {"rotatorRef": "r1", "date": "2026-05-04"}},
        ).json()
        assert again["ok"] is True
        assignments = again["state"]["inpatientAssignments"]
        assert len(assignments) == 2
        assert next(item for item in assignments if item["rotatorId"] == "r1")["source"] == "Manual"

        senior = client.post(
            "/api/scheduler/command",
            json={
                "type": "inpatient.assign",
                "input": {"rotatorRef": "r1", "date": "2026-05-04", "role": "Team Senior"},
            },
        ).json()
        assert senior["ok"] is True
        assignments = senior["state"]["inpatientAssignments"]
        assert len(assignments) == 3
        assert sorted(item["role"] for item in assignments if item["rotatorId"] == "r1") == ["Resident", "Team Senior"]

        after = client.get("/api/scheduler/state").json()
        assert len(after["inpatientAssignments"]) == 3


def test_inpatient_assign_rejects_invalid_active_date_before_persisting(tmp_path):
    with TestClient(create_app(data_dir=tmp_path)) as client:
        state = client.get("/api/scheduler/state").json()
        state.update(
            {
                "activeBlockId": "b1",
                "serviceBlocks": [
                    {
                        "id": "b1",
                        "name": "May 2026",
                        "startDate": "2026-05-04",
                        "endDate": "2026-05-31",
                        "status": "Draft",
                        "generate": {},
                        "holidays": [],
                    }
                ],
                "rotators": [
                    {
                        "id": "r1",
                        "fullName": "Drew Quinn",
                        "displayName": "Drew Quinn",
                        "program": "Other",
                        "level": "PGY-2",
                        "role": "Resident",
                        "segments": [{"start": "2026-05-04", "end": "2026-05-31"}],
                        "schoolType": "other",
                        "continuityClinic": "",
                        "dayOff": [],
                        "unavailableRanges": [{"start": "2026-05-07", "end": "2026-05-07", "label": "Vacation"}],
                    },
                    {
                        "id": "r2",
                        "fullName": "Blake Lee",
                        "displayName": "Blake Lee",
                        "program": "Other",
                        "level": "PGY-3",
                        "role": "Resident",
                        "segments": [{"start": "2026-06-01", "end": "2026-06-30"}],
                        "schoolType": "other",
                        "continuityClinic": "",
                        "dayOff": [],
                        "unavailableRanges": [],
                    },
                ],
                "inpatientAssignments": [],
                "outpatientSessions": [
                    {
                        "id": "op-2026-05-08-r1-AM",
                        "date": "2026-05-08",
                        "period": "AM",
                        "rotatorId": "r1",
                        "clinic": "Continuity",
                        "provider": "",
                        "source": "Manual",
                    }
                ],
            }
        )
        assert client.post("/api/scheduler/state", json=state).json()["ok"] is True

        ok = client.post(
            "/api/scheduler/command",
            json={"type": "inpatient.assign", "input": {"rotatorRef": "Drew Quinn", "date": "2026-05-06"}},
        ).json()
        assert ok["ok"] is True
        assert len(ok["state"]["inpatientAssignments"]) == 1

        for payload in [
            {"rotatorRef": "Blake Lee", "date": "2026-06-02"},
            {"rotatorRef": "Blake Lee", "date": "2026-05-06"},
            {"rotatorRef": "Drew Quinn", "date": "2026-05-07"},
            {"rotatorRef": "Drew Quinn", "date": "2026-05-08"},
        ]:
            blocked = client.post(
                "/api/scheduler/command",
                json={"type": "inpatient.assign", "input": payload},
            ).json()
            assert blocked["ok"] is False
            assert blocked["changed"] is False
            assert blocked["error"]["code"] == "invalid_inpatient_assignment"
            assert blocked["state"]["inpatientAssignments"] == ok["state"]["inpatientAssignments"]

        after = client.get("/api/scheduler/state").json()
        assert after["inpatientAssignments"] == ok["state"]["inpatientAssignments"]


def test_inpatient_drop_validates_and_persists_drag_drop(tmp_path):
    with TestClient(create_app(data_dir=tmp_path)) as client:
        state = client.get("/api/scheduler/state").json()
        state.update(
            {
                "rotators": [
                    {
                        "id": "r1",
                        "fullName": "Drew Quinn",
                        "displayName": "Drew Quinn",
                        "program": "Other",
                        "level": "PGY-2",
                        "role": "Resident",
                        "segments": [{"start": "2026-05-04", "end": "2026-05-08"}],
                        "schoolType": "other",
                        "continuityClinic": "",
                        "dayOff": [],
                        "unavailableRanges": [],
                    }
                ],
                "inpatientAssignments": [],
                "outpatientSessions": [],
            }
        )
        assert client.post("/api/scheduler/state", json=state).json()["ok"] is True

        dropped = client.post(
            "/api/scheduler/command",
            json={"type": "inpatient.drop", "input": {"rotatorRef": "Drew Quinn", "date": "2026-05-04"}},
        ).json()
        assert dropped["ok"] is True
        assert dropped["changed"] is True
        assert dropped["data"] == {"rotatorId": "r1", "date": "2026-05-04", "role": "Resident"}
        assert dropped["state"]["inpatientAssignments"] == [
            {
                "id": "in-2026-05-04-r1-resident",
                "date": "2026-05-04",
                "rotatorId": "r1",
                "role": "Resident",
                "source": "Drag-Drop",
            }
        ]

        blocked = client.post(
            "/api/scheduler/command",
            json={"type": "inpatient.drop", "input": {"rotatorRef": "Drew Quinn", "date": "2027-01-01"}},
        ).json()
        assert blocked["ok"] is False
        assert blocked["changed"] is False
        assert blocked["error"]["code"] == "invalid_drop"
        assert blocked["state"]["inpatientAssignments"] == dropped["state"]["inpatientAssignments"]


def test_inpatient_fellow_resolve_validates_before_persisting_batch(tmp_path):
    with TestClient(create_app(data_dir=tmp_path)) as client:
        state = client.get("/api/scheduler/state").json()
        state.update(
            {
                "rotators": [
                    {
                        "id": "f1",
                        "fullName": "Sam Carter",
                        "displayName": "Sam Carter",
                        "program": "Other",
                        "level": "Fellow",
                        "role": "Fellow",
                        "segments": [{"start": "2026-05-04", "end": "2026-05-31"}],
                        "schoolType": "other",
                        "continuityClinic": "Tuesday AM",
                        "dayOff": [],
                        "unavailableRanges": [],
                    },
                    {
                        "id": "r1",
                        "fullName": "Drew Quinn",
                        "displayName": "Drew Quinn",
                        "program": "Other",
                        "level": "PGY-2",
                        "role": "Resident",
                        "segments": [{"start": "2026-05-04", "end": "2026-05-31"}],
                        "schoolType": "other",
                        "continuityClinic": "",
                        "dayOff": [],
                        "unavailableRanges": [],
                    },
                ],
                "inpatientAssignments": [],
                "outpatientSessions": [],
            }
        )
        assert client.post("/api/scheduler/state", json=state).json()["ok"] is True

        resolved = client.post(
            "/api/scheduler/command",
            json={"type": "inpatient.fellow.resolve", "input": {"rotatorRef": "Sam Carter", "dates": ["2026-05-04", "2026-05-06"]}},
        ).json()
        assert resolved["ok"] is True
        assert resolved["data"] == {"rotatorId": "f1", "dates": ["2026-05-04", "2026-05-06"], "role": "Fellow"}
        assert [
            (item["date"], item["rotatorId"], item["role"], item["source"])
            for item in resolved["state"]["inpatientAssignments"]
        ] == [
            ("2026-05-04", "f1", "Fellow", "Manual"),
            ("2026-05-06", "f1", "Fellow", "Manual"),
        ]
        assert client.get("/api/scheduler/state").json()["inpatientAssignments"] == resolved["state"]["inpatientAssignments"]

        non_fellow = client.post(
            "/api/scheduler/command",
            json={"type": "inpatient.fellow.resolve", "input": {"rotatorRef": "Drew Quinn", "dates": ["2026-05-07"]}},
        ).json()
        assert non_fellow["ok"] is False
        assert non_fellow["error"]["code"] == "invalid_fellow_resolution"
        assert non_fellow["state"]["inpatientAssignments"] == resolved["state"]["inpatientAssignments"]

        blocked = client.post(
            "/api/scheduler/command",
            json={"type": "inpatient.fellow.resolve", "input": {"rotatorRef": "Sam Carter", "dates": ["2026-05-07", "2027-01-01"]}},
        ).json()
        assert blocked["ok"] is False
        assert blocked["changed"] is False
        assert blocked["error"]["code"] == "invalid_fellow_resolution"
        assert blocked["error"]["date"] == "2027-01-01"
        assert blocked["state"]["inpatientAssignments"] == resolved["state"]["inpatientAssignments"]
        assert client.get("/api/scheduler/state").json()["inpatientAssignments"] == resolved["state"]["inpatientAssignments"]


def test_assignment_delete_commands_remove_single_cells_and_persist(tmp_path):
    with TestClient(create_app(data_dir=tmp_path)) as client:
        state = client.get("/api/scheduler/state").json()
        state.update(
            {
                "inpatientAssignments": [
                    {
                        "id": "ip-1",
                        "date": "2026-05-04",
                        "rotatorId": "r1",
                        "role": "Resident",
                        "source": "Manual",
                    },
                    {
                        "id": "ip-2",
                        "date": "2026-05-05",
                        "rotatorId": "r2",
                        "role": "Resident",
                        "source": "Manual",
                    },
                ],
                "outpatientSessions": [
                    {
                        "id": "op-1",
                        "date": "2026-05-04",
                        "period": "AM",
                        "clinic": "Continuity",
                        "provider": "",
                        "rotatorId": "r1",
                        "status": "Scheduled",
                    },
                    {
                        "id": "op-2",
                        "date": "2026-05-05",
                        "period": "PM",
                        "clinic": "Continuity",
                        "provider": "",
                        "rotatorId": "r2",
                        "status": "Scheduled",
                    },
                ],
            }
        )
        assert client.post("/api/scheduler/state", json=state).json()["ok"] is True

        ip_deleted = client.post(
            "/api/scheduler/command",
            json={"type": "inpatient.delete", "input": {"assignmentRef": "ip-1"}},
        ).json()
        assert ip_deleted["ok"] is True
        assert ip_deleted["changed"] is True
        assert ip_deleted["data"] == {"assignmentId": "ip-1"}
        assert [item["id"] for item in ip_deleted["state"]["inpatientAssignments"]] == ["ip-2"]
        assert [item["id"] for item in ip_deleted["state"]["outpatientSessions"]] == ["op-1", "op-2"]

        op_deleted = client.post(
            "/api/scheduler/command",
            json={"type": "outpatient.delete", "input": {"sessionRef": "op-2"}},
        ).json()
        assert op_deleted["ok"] is True
        assert op_deleted["changed"] is True
        assert op_deleted["data"] == {"sessionId": "op-2"}
        assert [item["id"] for item in op_deleted["state"]["inpatientAssignments"]] == ["ip-2"]
        assert [item["id"] for item in op_deleted["state"]["outpatientSessions"]] == ["op-1"]

        missing = client.post(
            "/api/scheduler/command",
            json={"type": "inpatient.delete", "input": {"assignmentRef": "missing"}},
        ).json()
        assert missing["ok"] is False
        assert missing["error"]["code"] == "assignment_not_found"

        after = client.get("/api/scheduler/state").json()
        assert [item["id"] for item in after["inpatientAssignments"]] == ["ip-2"]
        assert [item["id"] for item in after["outpatientSessions"]] == ["op-1"]


def test_outpatient_assign_command_generates_and_persists(tmp_path):
    with TestClient(create_app(data_dir=tmp_path)) as client:
        state = client.get("/api/scheduler/state").json()
        state.update(
            {
                "activeBlockId": "b1",
                "serviceBlocks": [
                    {
                        "id": "b1",
                        "name": "May 2026",
                        "startDate": "2026-05-04",
                        "endDate": "2026-05-31",
                        "status": "Draft",
                        "generate": {},
                        "holidays": [],
                        "coverage": {},
                    }
                ],
                "rotators": [
                    {
                        "id": "r1",
                        "fullName": "Drew Quinn",
                        "displayName": "Drew Quinn",
                        "program": "Other",
                        "level": "PGY-2",
                        "role": "Resident",
                        "segments": [{"start": "2026-05-04", "end": "2026-05-08"}],
                        "schoolType": "other",
                        "continuityClinic": "",
                        "dayOff": [],
                        "unavailableRanges": [],
                    }
                ],
                "outpatientSessions": [],
            }
        )
        assert client.post("/api/scheduler/state", json=state).json()["ok"] is True

        response = client.post(
            "/api/scheduler/command",
            json={
                "type": "outpatient.assign",
                "input": {
                    "rotatorRef": "Drew Quinn",
                    "date": "2026-05-04",
                    "period": "AM",
                    "clinic": "QRS",
                    "provider": "Alder",
                    "details": [
                        {
                            "clinic": "Resident Continuity",
                            "attending": "Alder",
                            "task": "New visits",
                            "notes": "Room 3",
                        },
                        {"clinic": "", "attending": "", "task": "", "notes": ""},
                    ],
                },
            },
        )
        assert response.status_code == 200
        body = response.json()
        assert body["ok"] is True
        assert body["changed"] is True
        assert body["data"] == {"rotatorId": "r1", "date": "2026-05-04", "period": "AM"}
        assert body["state"]["outpatientSessions"] == [
            {
                "id": "out-2026-05-04-am-r1",
                "date": "2026-05-04",
                "period": "AM",
                "clinic": "QRS",
                "provider": "Alder",
                "rotatorId": "r1",
                "status": "Scheduled",
                "details": [
                    {
                        "clinic": "Resident Continuity",
                        "attending": "Alder",
                        "task": "New visits",
                        "notes": "Room 3",
                    }
                ],
            }
        ]

        after = client.get("/api/scheduler/state").json()
        assert len(after["outpatientSessions"]) == 1


def test_outpatient_assign_rejects_invalid_period(tmp_path):
    with TestClient(create_app(data_dir=tmp_path)) as client:
        response = client.post(
            "/api/scheduler/command",
            json={
                "type": "outpatient.assign",
                "input": {"rotatorRef": "missing", "date": "2026-05-04", "period": "Noon"},
            },
        )
        body = response.json()
        assert response.status_code == 200
        assert body["ok"] is False
        assert body["error"]["code"] == "invalid_enum"


def test_outpatient_assign_rejects_closed_or_ineligible_dates(tmp_path):
    with TestClient(create_app(data_dir=tmp_path)) as client:
        state = client.get("/api/scheduler/state").json()
        state.update(
            {
                "activeBlockId": "b1",
                "serviceBlocks": [
                    {
                        "id": "b1",
                        "name": "May 2026",
                        "startDate": "2026-05-04",
                        "endDate": "2026-05-31",
                        "status": "Draft",
                        "generate": {},
                        "holidays": [{"date": "2026-05-25", "noClinic": True, "label": "Memorial Day"}],
                        "coverage": {},
                    }
                ],
                "rotators": [
                    {
                        "id": "r1",
                        "fullName": "Drew Quinn",
                        "displayName": "Drew Quinn",
                        "program": "Other",
                        "level": "PGY-2",
                        "role": "Resident",
                        "segments": [{"start": "2026-05-04", "end": "2026-05-31"}],
                        "schoolType": "other",
                        "continuityClinic": "",
                        "dayOff": [],
                        "unavailableRanges": [{"start": "2026-05-07", "end": "2026-05-07", "label": "Vacation"}],
                    }
                ],
                "outpatientSessions": [],
            }
        )
        assert client.post("/api/scheduler/state", json=state).json()["ok"] is True

        valid = client.post(
            "/api/scheduler/command",
            json={"type": "outpatient.assign", "input": {"rotatorRef": "Drew Quinn", "date": "2026-05-06", "period": "AM"}},
        ).json()
        assert valid["ok"] is True
        assert len(valid["state"]["outpatientSessions"]) == 1

        for date in ["2026-05-09", "2026-05-25", "2026-05-07", "2026-06-01"]:
            blocked = client.post(
                "/api/scheduler/command",
                json={"type": "outpatient.assign", "input": {"rotatorRef": "Drew Quinn", "date": date, "period": "PM"}},
            ).json()
            assert blocked["ok"] is False
            assert blocked["changed"] is False
            assert blocked["error"]["code"] == "invalid_outpatient_assignment"
            assert blocked["error"]["date"] == date
            assert blocked["state"]["outpatientSessions"] == valid["state"]["outpatientSessions"]
            assert client.get("/api/scheduler/state").json()["outpatientSessions"] == valid["state"]["outpatientSessions"]


def test_clinic_assign_command_validates_and_persists(tmp_path):
    with TestClient(create_app(data_dir=tmp_path)) as client:
        state = client.get("/api/scheduler/state").json()
        block = {
            "id": "b1",
            "name": "Clinic Week",
            "startDate": "2026-05-04",
            "endDate": "2026-05-08",
            "status": "Draft",
            "generate": {},
            "holidays": [],
            "coverage": {"weekday": {"ip": {"count": 1}}},
        }
        state.update(
            {
                "activeBlockId": "b1",
                "serviceBlocks": [block],
                "rotators": [
                    {
                        "id": "r1",
                        "fullName": "Drew Quinn",
                        "displayName": "Drew Quinn",
                        "program": "Other",
                        "level": "PGY-2",
                        "role": "Resident",
                        "segments": [{"start": "2026-05-04", "end": "2026-05-08"}],
                        "schoolType": "other",
                        "continuityClinic": "",
                        "dayOff": [],
                        "unavailableRanges": [],
                    }
                ],
                "attendings": [
                    {
                        "name": "Alder",
                        "recurringClinics": [
                            {"weekday": "Monday", "session": "AM", "clinicName": "General Neuro", "capacity": 1}
                        ],
                        "oneOffDates": [],
                    }
                ],
                "outpatientSessions": [
                    {"id": "op-1", "date": "2026-05-04", "period": "AM", "clinic": "", "provider": "", "rotatorId": "r1", "status": "Scheduled"}
                ],
                "clinicAssignments": [],
            }
        )
        assert client.post("/api/scheduler/state", json=state).json()["ok"] is True

        occurrence_id = "clinic-occurrence::alder::0-monday-AM-general-neuro::2026-05-04::AM"
        assigned = client.post(
            "/api/scheduler/command",
            json={
                "type": "clinic.assign",
                "input": {
                    "clinicOccurrenceId": occurrence_id,
                    "rotatorRef": "Drew Quinn",
                    "date": "2026-05-04",
                    "session": "AM",
                },
            },
        ).json()
        assert assigned["ok"] is True
        assert assigned["changed"] is True
        assignment_id = assigned["data"]["assignmentId"]
        assert assigned["state"]["clinicAssignments"] == [
            {
                "id": assignment_id,
                "clinicOccurrenceId": occurrence_id,
                "rotatorId": "r1",
                "date": "2026-05-04",
                "session": "AM",
                "source": "manual",
            }
        ]
        after_assign = client.get("/api/scheduler/state").json()
        assert after_assign["clinicAssignments"] == assigned["state"]["clinicAssignments"]

        # Idempotent re-assign stays a single persisted row.
        again = client.post(
            "/api/scheduler/command",
            json={
                "type": "clinic.assign",
                "input": {
                    "clinicOccurrenceId": occurrence_id,
                    "rotatorRef": "r1",
                    "date": "2026-05-04",
                    "session": "AM",
                },
            },
        ).json()
        assert again["ok"] is True
        assert len(again["state"]["clinicAssignments"]) == 1

        deleted = client.post(
            "/api/scheduler/command",
            json={"type": "clinic.delete", "input": {"assignmentRef": assignment_id}},
        ).json()
        assert deleted["ok"] is True
        assert deleted["state"]["clinicAssignments"] == []
        assert deleted["state"]["outpatientSessions"] == state["outpatientSessions"]

        after = client.get("/api/scheduler/state").json()
        assert after["clinicAssignments"] == []


def test_clinic_assign_rejects_non_outpatient_and_same_session_double_book(tmp_path):
    with TestClient(create_app(data_dir=tmp_path)) as client:
        state = client.get("/api/scheduler/state").json()
        state.update(
            {
                "activeBlockId": "b1",
                "serviceBlocks": [
                    {
                        "id": "b1",
                        "name": "Clinic Week",
                        "startDate": "2026-05-04",
                        "endDate": "2026-05-08",
                        "status": "Draft",
                        "generate": {},
                        "holidays": [],
                        "coverage": {"weekday": {"ip": {"count": 1}}},
                    }
                ],
                "rotators": [
                    {
                        "id": "r1",
                        "fullName": "Drew Quinn",
                        "displayName": "Drew Quinn",
                        "program": "Other",
                        "level": "PGY-2",
                        "role": "Resident",
                        "segments": [{"start": "2026-05-04", "end": "2026-05-08"}],
                        "schoolType": "other",
                        "continuityClinic": "",
                        "dayOff": [],
                        "unavailableRanges": [],
                    },
                    {
                        "id": "r2",
                        "fullName": "Bea Student",
                        "displayName": "Bea Student",
                        "program": "UT Med Student",
                        "level": "MS3",
                        "role": "Student",
                        "segments": [{"start": "2026-05-04", "end": "2026-05-08"}],
                        "schoolType": "ut-student",
                        "continuityClinic": "",
                        "dayOff": [],
                        "unavailableRanges": [],
                    },
                ],
                "attendings": [
                    {
                        "name": "Alder",
                        "recurringClinics": [
                            {"weekday": "Monday", "session": "AM", "clinicName": "General Neuro", "capacity": 2, "allowedRoles": ["Resident"]},
                            {"weekday": "Monday", "session": "AM", "clinicName": "Epilepsy", "capacity": 2},
                        ],
                        "oneOffDates": [],
                    }
                ],
                "outpatientSessions": [],
                "clinicAssignments": [],
            }
        )
        assert client.post("/api/scheduler/state", json=state).json()["ok"] is True

        general = "clinic-occurrence::alder::0-monday-AM-general-neuro::2026-05-04::AM"
        not_op = client.post(
            "/api/scheduler/command",
            json={"type": "clinic.assign", "input": {"clinicOccurrenceId": general, "rotatorRef": "r1", "date": "2026-05-04", "session": "AM"}},
        ).json()
        assert not_op["ok"] is False
        assert not_op["error"]["code"] == "not-outpatient"
        assert not_op["state"]["clinicAssignments"] == []

        state["outpatientSessions"] = [
            {"id": "op-1", "date": "2026-05-04", "period": "AM", "clinic": "", "provider": "", "rotatorId": "r1", "status": "Scheduled"},
            {"id": "op-2", "date": "2026-05-04", "period": "AM", "clinic": "", "provider": "", "rotatorId": "r2", "status": "Scheduled"},
        ]
        assert client.post("/api/scheduler/state", json=state).json()["ok"] is True
        role_denied = client.post(
            "/api/scheduler/command",
            json={"type": "clinic.assign", "input": {"clinicOccurrenceId": general, "rotatorRef": "r2", "date": "2026-05-04", "session": "AM"}},
        ).json()
        assert role_denied["ok"] is False
        assert role_denied["error"]["code"] == "role-not-allowed"
        assert role_denied["state"]["clinicAssignments"] == []

        assigned = client.post(
            "/api/scheduler/command",
            json={"type": "clinic.assign", "input": {"clinicOccurrenceId": general, "rotatorRef": "r1", "date": "2026-05-04", "session": "AM"}},
        ).json()
        assert assigned["ok"] is True

        epilepsy = "clinic-occurrence::alder::1-monday-AM-epilepsy::2026-05-04::AM"
        duplicate = client.post(
            "/api/scheduler/command",
            json={"type": "clinic.assign", "input": {"clinicOccurrenceId": epilepsy, "rotatorRef": "r1", "date": "2026-05-04", "session": "AM"}},
        ).json()
        assert duplicate["ok"] is False
        assert duplicate["error"]["code"] == "already-in-session"
        assert len(duplicate["state"]["clinicAssignments"]) == 1


def test_daily_report_command_returns_report_without_mutation(tmp_path):
    with TestClient(create_app(data_dir=tmp_path)) as client:
        state = client.get("/api/scheduler/state").json()
        block = {
            "id": "b1",
            "name": "Report Week",
            "startDate": "2026-05-04",
            "endDate": "2026-05-08",
            "status": "Draft",
            "generate": {},
            "holidays": [],
            "coverage": {"weekday": {"ip": {"count": 1}}},
        }
        state.update(
            {
                "activeBlockId": "b1",
                "serviceBlocks": [block],
                "rotators": [
                    {
                        "id": "r1",
                        "fullName": "Drew Quinn",
                        "displayName": "Drew Quinn",
                        "program": "Other",
                        "level": "PGY-2",
                        "role": "Resident",
                        "segments": [{"start": "2026-05-04", "end": "2026-05-08"}],
                        "schoolType": "other",
                        "continuityClinic": "",
                        "dayOff": [],
                        "unavailableRanges": [],
                    }
                ],
                "inpatientAssignments": [
                    {
                        "id": "in-2026-05-04-r1",
                        "date": "2026-05-04",
                        "rotatorId": "r1",
                        "role": "Resident",
                        "source": "Manual",
                    },
                    {
                        "id": "in-off-2026-05-05-r1",
                        "date": "2026-05-05",
                        "rotatorId": "r1",
                        "role": "Off",
                        "source": "Range-Assigned",
                    }
                ],
                "outpatientSessions": [
                    {
                        "id": "out-2026-05-04-am-r1",
                        "date": "2026-05-04",
                        "period": "AM",
                        "clinic": "QRS",
                        "provider": "Alder",
                        "rotatorId": "r1",
                        "status": "Scheduled",
                    }
                ],
            }
        )
        assert client.post("/api/scheduler/state", json=state).json()["ok"] is True

        response = client.post(
            "/api/scheduler/command",
            json={"type": "report.daily", "input": {"date": "2026-05-04"}},
        )
        assert response.status_code == 200
        body = response.json()
        assert body["ok"] is True
        assert body["changed"] is False
        assert body["message"] == "Generated daily report for 2026-05-04."
        assert body["data"]["report"].splitlines()[:5] == [
            "Daily Team Report",
            "Report Week",
            "2026-05-04",
            "",
            "Inpatient",
        ]
        assert "Resident: Drew Quinn" in body["data"]["report"]
        assert "AM QRS: Drew Quinn" in body["data"]["report"]
        assert "Drew Quinn is double-booked" in body["data"]["report"]


def test_conflicts_list_command_returns_full_and_filtered_lists(tmp_path):
    with TestClient(create_app(data_dir=tmp_path)) as client:
        state = client.get("/api/scheduler/state").json()
        block = {
            "id": "b1",
            "name": "Conflict Week",
            "startDate": "2026-05-04",
            "endDate": "2026-05-05",
            "status": "Draft",
            "generate": {},
            "holidays": [],
            "coverage": {"weekday": {"ip": {"count": 1}}},
        }
        state.update(
            {
                "activeBlockId": "b1",
                "serviceBlocks": [block],
                "rotators": [
                    {
                        "id": "r1",
                        "fullName": "Drew Quinn",
                        "displayName": "Drew Quinn",
                        "program": "Other",
                        "level": "PGY-2",
                        "role": "Resident",
                        "segments": [{"start": "2026-05-04", "end": "2026-05-05"}],
                        "schoolType": "other",
                        "continuityClinic": "",
                        "dayOff": [],
                        "unavailableRanges": [],
                    }
                ],
                "inpatientAssignments": [
                    {
                        "id": "in-2026-05-04-r1",
                        "date": "2026-05-04",
                        "rotatorId": "r1",
                        "role": "Resident",
                        "source": "Manual",
                    }
                ],
                "outpatientSessions": [
                    {
                        "id": "out-2026-05-04-am-r1",
                        "date": "2026-05-04",
                        "period": "AM",
                        "clinic": "QRS",
                        "provider": "Alder",
                        "rotatorId": "r1",
                        "status": "Scheduled",
                    }
                ],
            }
        )
        assert client.post("/api/scheduler/state", json=state).json()["ok"] is True

        all_conflicts = client.post("/api/scheduler/command", json={"type": "conflicts.list"}).json()
        assert all_conflicts["ok"] is True
        assert all_conflicts["changed"] is False
        assert all_conflicts["message"] == "Found 2 conflicts."
        assert [conflict["title"] for conflict in all_conflicts["data"]["conflicts"]] == [
            "Drew Quinn is double-booked",
            "No inpatient coverage on 2026-05-05",
        ]

        filtered = client.post(
            "/api/scheduler/command",
            json={"type": "conflicts.list", "input": {"date": "2026-05-04"}},
        ).json()
        assert filtered["ok"] is True
        assert filtered["message"] == "Found 1 conflicts."
        assert filtered["data"]["conflicts"][0]["date"] == "2026-05-04"


def test_conflicts_list_pins_individual_post_hoc_conflict_types(tmp_path):
    with TestClient(create_app(data_dir=tmp_path)) as client:
        state = client.get("/api/scheduler/state").json()
        block = {
            "id": "b1",
            "name": "Conflict Coverage",
            "startDate": "2026-05-04",
            "endDate": "2026-05-10",
            "status": "Draft",
            "generate": {},
            "holidays": [{"date": "2026-05-06", "label": "No Clinic Day", "noClinic": True}],
            "coverage": {
                "weekday": {"ip": {"count": 0}},
                "saturday": {"ip": {"count": 0}},
                "sunday": {"ip": {"count": 0}},
                "holiday": {"ip": {"count": 0}},
            },
        }
        state.update(
            {
                "activeBlockId": "b1",
                "serviceBlocks": [block],
                "rotators": [
                    {
                        "id": "r1",
                        "fullName": "Drew Quinn",
                        "displayName": "Drew Quinn",
                        "program": "Other",
                        "level": "PGY-2",
                        "role": "Resident",
                        "segments": [{"start": "2026-05-04", "end": "2026-05-10"}],
                        "schoolType": "other",
                        "continuityClinic": "Tuesday PM",
                        "dayOff": ["Monday"],
                        "unavailableRanges": [{"start": "2026-05-07", "end": "2026-05-07", "label": "Conference"}],
                    },
                    {
                        "id": "r-phantom",
                        "fullName": "Phantom Resident",
                        "displayName": "Phantom Resident",
                        "program": "Other",
                        "level": "PGY-2",
                        "role": "Resident",
                        "segments": [],
                        "schoolType": "other",
                        "continuityClinic": "",
                        "dayOff": [],
                        "unavailableRanges": [],
                    },
                ],
                "inpatientAssignments": [
                    {"id": "in-day-off", "date": "2026-05-04", "rotatorId": "r1", "role": "Resident", "source": "Manual"},
                    {"id": "in-continuity", "date": "2026-05-05", "rotatorId": "r1", "role": "Resident", "source": "Manual"},
                    {"id": "in-phantom", "date": "2026-05-08", "rotatorId": "r-phantom", "role": "Resident", "source": "Manual"},
                ],
                "outpatientSessions": [
                    {"id": "op-continuity", "date": "2026-05-05", "period": "PM", "clinic": "QRS", "provider": "Alder", "rotatorId": "r1", "status": "Scheduled"},
                    {"id": "op-holiday", "date": "2026-05-06", "period": "AM", "clinic": "QRS", "provider": "Alder", "rotatorId": "r1", "status": "Scheduled"},
                    {"id": "op-unavailable", "date": "2026-05-07", "period": "AM", "clinic": "QRS", "provider": "Alder", "rotatorId": "r1", "status": "Scheduled"},
                    {"id": "op-weekend", "date": "2026-05-09", "period": "AM", "clinic": "QRS", "provider": "Alder", "rotatorId": "r1", "status": "Scheduled"},
                ],
            }
        )
        assert client.post("/api/scheduler/state", json=state).json()["ok"] is True

        body = client.post("/api/scheduler/command", json={"type": "conflicts.list"}).json()
        assert body["ok"] is True
        conflicts = body["data"]["conflicts"]
        by_type = {}
        for conflict in conflicts:
            by_type.setdefault(conflict["type"], []).append(conflict)

        unavailable = by_type["rotator-unavailable"]
        assert any(conflict["assignment"] == "inpatient" and conflict["date"] == "2026-05-04" for conflict in unavailable)
        assert any(conflict["assignment"] == "outpatient" and conflict["date"] == "2026-05-07" for conflict in unavailable)
        assert by_type["holiday-clinic"][0]["date"] == "2026-05-06"
        assert by_type["outpatient-weekend"][0]["date"] == "2026-05-09"
        assert any(conflict["assignment"] == "inpatient" and conflict["date"] == "2026-05-05" for conflict in by_type["continuity-clinic-conflict"])
        assert any(conflict["assignment"] == "outpatient" and conflict["period"] == "PM" for conflict in by_type["continuity-clinic-conflict"])
        assert by_type["missing-legend"][0]["rotatorId"] == "r-phantom"


def test_conflicts_list_includes_clinic_assignment_conflicts(tmp_path):
    with TestClient(create_app(data_dir=tmp_path)) as client:
        state = client.get("/api/scheduler/state").json()
        block = {
            "id": "b1",
            "name": "Clinic Conflicts",
            "startDate": "2026-05-04",
            "endDate": "2026-05-04",
            "status": "Draft",
            "generate": {},
            "holidays": [],
            "coverage": {"weekday": {"ip": {"count": 0}}},
        }
        state.update(
            {
                "activeBlockId": "b1",
                "serviceBlocks": [block],
                "rotators": [
                    {
                        "id": "r1",
                        "fullName": "Drew Quinn",
                        "displayName": "Drew Quinn",
                        "program": "Other",
                        "level": "PGY-2",
                        "role": "Resident",
                        "segments": [{"start": "2026-05-04", "end": "2026-05-04"}],
                        "schoolType": "other",
                        "continuityClinic": "",
                        "dayOff": [],
                        "unavailableRanges": [],
                    },
                    {
                        "id": "r2",
                        "fullName": "Bea Lee",
                        "displayName": "Bea Lee",
                        "program": "Other",
                        "level": "PGY-3",
                        "role": "Resident",
                        "segments": [{"start": "2026-05-04", "end": "2026-05-04"}],
                        "schoolType": "other",
                        "continuityClinic": "",
                        "dayOff": [],
                        "unavailableRanges": [],
                    },
                ],
                "attendings": [
                    {
                        "name": "Alder",
                        "recurringClinics": [
                            {"weekday": "Monday", "session": "AM", "clinicName": "General Neuro", "capacity": 1},
                            {"weekday": "Monday", "session": "AM", "clinicName": "Epilepsy", "capacity": 1},
                        ],
                        "oneOffDates": [],
                    }
                ],
                "outpatientSessions": [
                    {"id": "op-r1", "date": "2026-05-04", "period": "AM", "clinic": "", "provider": "", "rotatorId": "r1", "status": "Scheduled"},
                    {"id": "op-r2", "date": "2026-05-04", "period": "AM", "clinic": "", "provider": "", "rotatorId": "r2", "status": "Scheduled"},
                ],
                "clinicAssignments": [
                    {
                        "id": "a1",
                        "clinicOccurrenceId": "clinic-occurrence::alder::0-monday-AM-general-neuro::2026-05-04::AM",
                        "rotatorId": "r1",
                        "date": "2026-05-04",
                        "session": "AM",
                        "source": "manual",
                    },
                    {
                        "id": "a2",
                        "clinicOccurrenceId": "clinic-occurrence::alder::1-monday-AM-epilepsy::2026-05-04::AM",
                        "rotatorId": "r1",
                        "date": "2026-05-04",
                        "session": "AM",
                        "source": "manual",
                    },
                    {
                        "id": "a3",
                        "clinicOccurrenceId": "clinic-occurrence::alder::0-monday-AM-general-neuro::2026-05-04::AM",
                        "rotatorId": "r2",
                        "date": "2026-05-04",
                        "session": "AM",
                        "source": "manual",
                    },
                ],
            }
        )
        assert client.post("/api/scheduler/state", json=state).json()["ok"] is True

        body = client.post("/api/scheduler/command", json={"type": "conflicts.list"}).json()
        assert body["ok"] is True
        clinic_conflicts = [
            conflict for conflict in body["data"]["conflicts"]
            if conflict["assignment"] == "clinic"
        ]
        assert [conflict["type"] for conflict in clinic_conflicts] == [
            "clinic-double-book",
            "clinic-over-capacity",
        ]
        assert clinic_conflicts[0]["title"] == "Drew Quinn is double-booked in clinic"
        assert clinic_conflicts[1]["title"] == "Clinic over capacity on 2026-05-04"


def test_conflicts_list_includes_clinic_role_policy_conflict(tmp_path):
    with TestClient(create_app(data_dir=tmp_path)) as client:
        state = client.get("/api/scheduler/state").json()
        block = {
            "id": "b1",
            "name": "Clinic Policy",
            "startDate": "2026-05-04",
            "endDate": "2026-05-04",
            "status": "Draft",
            "generate": {},
            "holidays": [],
            "coverage": {"weekday": {"ip": {"count": 0}}},
        }
        occurrence_id = "clinic-occurrence::alder::0-monday-AM-general-neuro::2026-05-04::AM"
        state.update(
            {
                "activeBlockId": "b1",
                "serviceBlocks": [block],
                "rotators": [
                    {
                        "id": "r2",
                        "fullName": "Bea Student",
                        "displayName": "Bea Student",
                        "program": "UT Med Student",
                        "level": "MS3",
                        "role": "Student",
                        "segments": [{"start": "2026-05-04", "end": "2026-05-04"}],
                        "schoolType": "ut-student",
                        "continuityClinic": "",
                        "dayOff": [],
                        "unavailableRanges": [],
                    }
                ],
                "attendings": [
                    {
                        "name": "Alder",
                        "recurringClinics": [
                            {"weekday": "Monday", "session": "AM", "clinicName": "General Neuro", "allowedRoles": ["Resident"]},
                        ],
                        "oneOffDates": [],
                    }
                ],
                "outpatientSessions": [
                    {"id": "op-r2", "date": "2026-05-04", "period": "AM", "clinic": "", "provider": "", "rotatorId": "r2", "status": "Scheduled"},
                ],
                "clinicAssignments": [
                    {
                        "id": "a-role",
                        "clinicOccurrenceId": occurrence_id,
                        "rotatorId": "r2",
                        "date": "2026-05-04",
                        "session": "AM",
                        "source": "manual",
                    }
                ],
            }
        )
        assert client.post("/api/scheduler/state", json=state).json()["ok"] is True

        body = client.post("/api/scheduler/command", json={"type": "conflicts.list"}).json()
        assert body["ok"] is True
        role_conflict = next(
            conflict for conflict in body["data"]["conflicts"]
            if conflict["type"] == "clinic-role-mismatch"
        )
        assert role_conflict["title"] == "Bea Student does not match clinic role policy"
        assert "Student" in role_conflict["detail"]
        assert "Resident" in role_conflict["detail"]


def test_export_package_command_returns_package_without_mutation(tmp_path):
    with TestClient(create_app(data_dir=tmp_path)) as client:
        state = client.get("/api/scheduler/state").json()
        block = {
            "id": "b1",
            "name": "Export Week",
            "startDate": "2026-05-04",
            "endDate": "2026-05-05",
            "status": "Draft",
            "generate": {},
            "holidays": [],
            "coverage": {"weekday": {"ip": {"count": 1}}},
        }
        state.update(
            {
                "activeBlockId": "b1",
                "serviceBlocks": [block],
                "rotators": [
                    {
                        "id": "r1",
                        "fullName": "Drew Quinn",
                        "displayName": "Drew Quinn",
                        "program": "Other",
                        "level": "PGY-2",
                        "role": "Resident",
                        "segments": [{"start": "2026-05-04", "end": "2026-05-05"}],
                        "schoolType": "other",
                        "continuityClinic": "",
                        "dayOff": [],
                        "unavailableRanges": [],
                    }
                ],
                "inpatientAssignments": [
                    {
                        "id": "in-2026-05-04-r1",
                        "date": "2026-05-04",
                        "rotatorId": "r1",
                        "role": "Resident",
                        "source": "Manual",
                    }
                ],
                "attendings": [
                    {
                        "name": "Alder",
                        "recurringClinics": [
                            {"weekday": "Monday", "session": "AM", "clinicName": "General Neuro", "capacity": 1}
                        ],
                        "oneOffDates": [],
                    }
                ],
                "outpatientSessions": [
                    {"id": "op-1", "date": "2026-05-04", "period": "AM", "clinic": "", "provider": "", "rotatorId": "r1", "status": "Scheduled"}
                ],
                "clinicAssignments": [
                    {
                        "id": "clinic-1",
                        "clinicOccurrenceId": "clinic-occurrence::alder::0-monday-AM-general-neuro::2026-05-04::AM",
                        "rotatorId": "r1",
                        "date": "2026-05-04",
                        "session": "AM",
                        "source": "manual",
                    }
                ],
                "expectedSourcePrograms": ["Other", "Missing Program"],
                "sources": [
                    {
                        "id": "source-1",
                        "fileName": "roster.csv",
                        "fileType": "csv",
                        "program": "Other",
                        "status": "Reviewed",
                        "importedAt": "2026-07-05",
                        "importedRotatorCount": 1,
                        "importWarnings": ["Row 2: missing level."],
                        "parsedRows": [{"Name": "Drew Quinn"}],
                    }
                ],
            }
        )
        assert client.post("/api/scheduler/state", json=state).json()["ok"] is True

        response = client.post("/api/scheduler/command", json={"type": "export.package"})
        assert response.status_code == 200
        body = response.json()
        assert body["ok"] is True
        assert body["changed"] is False
        package = body["data"]["package"]
        assert body["message"] == "Built export package."
        assert [file["name"] for file in package["manifest"]["files"]] == [
            "manifest.json",
            "roster.json",
            "roster.csv",
            "inpatient-calendar.json",
            "inpatient-calendar.csv",
            "outpatient-calendar.json",
            "outpatient-calendar.csv",
            "daily-reports.txt",
            "legend.json",
            "legend.csv",
            "conflicts.json",
            "conflicts.csv",
            "source-import-summary.json",
            "source-import-summary.csv",
            "schedule-package.json",
        ]
        assert [file["key"] for file in package["manifest"]["files"]] == [
            "manifest",
            "roster",
            "rosterCsv",
            "inpatientCalendar",
            "inpatientCalendarCsv",
            "outpatientCalendar",
            "outpatientCalendarCsv",
            "dailyReports",
            "legend",
            "legendCsv",
            "conflicts",
            "conflictsCsv",
            "sourceSummary",
            "sourceSummaryCsv",
            "schedulePackage",
        ]
        assert package["manifest"]["block"] == "Export Week"
        assert package["manifest"]["exportVersion"] == 2
        files_by_key = {file["key"]: file for file in package["manifest"]["files"]}
        assert files_by_key["rosterCsv"]["format"] == "csv"
        assert files_by_key["rosterCsv"]["rows"] == 1
        assert files_by_key["inpatientCalendarCsv"]["rows"] == 2
        assert files_by_key["outpatientCalendarCsv"]["rows"] == 3
        assert package["inpatientCalendar"][0]["date"] == "2026-05-04"
        assert package["inpatientCalendar"][0]["assignments"][0]["rotatorName"] == "Drew Quinn"
        assert all(
            assignment.get("role") != "Off"
            for day in package["inpatientCalendar"]
            for assignment in day["assignments"]
        )
        assert package["outpatientCalendar"][0]["sessions"][0]["rotatorName"] == "Drew Quinn"
        assert package["outpatientCalendar"][0]["clinicAssignments"][0]["clinicName"] == "General Neuro"
        assert package["legend"]["entries"][0]["displayLabel"] == "Drew Quinn"
        assert package["sourceSummary"]["expectedPrograms"] == [
            {"program": "Other", "status": "present"},
            {"program": "Missing Program", "status": "missing"},
        ]
        assert package["sourceSummary"]["sources"][0]["fileName"] == "roster.csv"
        assert package["sourceSummary"]["sources"][0]["importWarningCount"] == 1
        assert "fullName,displayName" in package["rosterCsv"]
        assert "Drew Quinn" in package["rosterCsv"]
        assert "date,assignmentId,rotatorId,rotatorName,role,source" in package["inpatientCalendarCsv"]
        assert "2026-05-05" in package["inpatientCalendarCsv"]
        assert "in-off-2026-05-05-r1" not in package["inpatientCalendarCsv"]
        assert "kind,sessionId,assignmentId" in package["outpatientCalendarCsv"]
        assert "General Neuro" in package["outpatientCalendarCsv"]
        assert "displayLabel,dateRange" in package["legendCsv"]
        assert "id,type,severity,title,detail,date,rotatorId,status,assignment,period" in package["conflictsCsv"]
        assert "double-booked" in package["conflictsCsv"]
        assert "Drew Quinn is double-booked" in package["conflictsCsv"]
        assert "source-import-summary.csv" in [file["name"] for file in package["manifest"]["files"]]
        assert "kind,id,fileName,fileType,program,status,importedAt,importedRotatorCount,rows,importWarningCount" in package["sourceSummaryCsv"]
        assert "roster.csv,csv,Other,Reviewed,2026-07-05,1,1,1" in package["sourceSummaryCsv"]
        assert "expectedProgram" in package["sourceSummaryCsv"]
        assert "Missing Program" in package["sourceSummaryCsv"]
        assert "Daily Team Report" in package["dailyReports"]
        assert "Clinic Assignments" in package["dailyReports"]
        assert "AM General Neuro with Alder: Drew Quinn" in package["dailyReports"]
        assert package["schedulePackage"]["activeBlockId"] == "b1"


def test_export_pdfs_command_returns_pdf_files_without_mutation(tmp_path):
    with TestClient(create_app(data_dir=tmp_path)) as client:
        state = client.get("/api/scheduler/state").json()
        block = {
            "id": "b1",
            "name": "PDF Week",
            "startDate": "2026-05-04",
            "endDate": "2026-05-05",
            "status": "Draft",
            "generate": {},
            "holidays": [],
            "coverage": {"weekday": {"ip": {"count": 1}}},
        }
        state.update(
            {
                "activeBlockId": "b1",
                "serviceBlocks": [block],
                "rotators": [
                    {
                        "id": "r1",
                        "fullName": "Drew Quinn",
                        "displayName": "Drew Quinn",
                        "program": "Other",
                        "level": "PGY-2",
                        "role": "Resident",
                        "segments": [{"start": "2026-05-04", "end": "2026-05-05"}],
                        "schoolType": "other",
                        "continuityClinic": "",
                        "dayOff": [],
                        "unavailableRanges": [],
                    }
                ],
                "inpatientAssignments": [
                    {
                        "id": "in-2026-05-04-r1",
                        "date": "2026-05-04",
                        "rotatorId": "r1",
                        "role": "Resident",
                        "source": "Manual",
                    }
                ],
                "outpatientSessions": [
                    {
                        "id": "out-2026-05-04-am-r1",
                        "date": "2026-05-04",
                        "period": "AM",
                        "clinic": "",
                        "provider": "",
                        "rotatorId": "r1",
                        "status": "Scheduled",
                    }
                ],
                "attendings": [
                    {
                        "name": "Alder",
                        "recurringClinics": [
                            {"weekday": "Monday", "session": "AM", "clinicName": "General Neuro", "capacity": 1}
                        ],
                        "oneOffDates": [],
                    }
                ],
                "clinicAssignments": [
                    {
                        "id": "clinic-1",
                        "clinicOccurrenceId": "clinic-occurrence::alder::0-monday-AM-general-neuro::2026-05-04::AM",
                        "rotatorId": "r1",
                        "date": "2026-05-04",
                        "session": "AM",
                        "source": "manual",
                    }
                ],
                "posterSettings": {
                    "programName": "Child Neuro Rotation",
                    "chief": "Dr. Lane",
                    "notes": ["Bring badge."],
                    "locations": [{"name": "North Clinic", "address": "123 Lane"}],
                    "tagline": "Thanks team.",
                },
            }
        )
        assert client.post("/api/scheduler/state", json=state).json()["ok"] is True

        response = client.post("/api/scheduler/command", json={"type": "export.pdfs"})
        assert response.status_code == 200
        body = response.json()
        assert body["ok"] is True
        assert body["changed"] is False
        files = body["data"]["files"]
        assert body["message"] == "Built 4 PDF exports."
        assert [file["name"] for file in files] == [
            "pdf-week-schedule.pdf",
            "pdf-week-inpatient-poster.pdf",
            "pdf-week-outpatient-poster.pdf",
            "pdf-week-outpatient-week-grid.pdf",
        ]
        for file in files:
            decoded = base64.b64decode(file["base64"])
            assert file["mimeType"] == "application/pdf"
            assert decoded.startswith(b"%PDF-1.4")
            assert decoded.rstrip().endswith(b"%%EOF")
            assert len(decoded) > 700
            text = decoded.decode("latin-1")
            if file["name"].endswith("-poster.pdf"):
                assert "Child Neuro Rotation" in text
                assert "Chief: Dr. Lane" in text
                assert "North Clinic - 123 Lane" in text
                assert "Bring badge." in text
                assert "Thanks team." in text
            if file["name"] in {"pdf-week-schedule.pdf", "pdf-week-outpatient-poster.pdf"}:
                assert "General Neuro with Alder" in text
                assert "Drew Quinn" in text


def test_export_word_command_returns_docx_files_without_mutation(tmp_path):
    with TestClient(create_app(data_dir=tmp_path)) as client:
        state = client.get("/api/scheduler/state").json()
        block = {
            "id": "b1",
            "name": "Word Week",
            "startDate": "2026-05-04",
            "endDate": "2026-05-05",
            "status": "Draft",
            "generate": {},
            "holidays": [],
            "coverage": {"weekday": {"ip": {"count": 1}}},
        }
        state.update(
            {
                "activeBlockId": "b1",
                "serviceBlocks": [block],
                "rotators": [
                    {
                        "id": "r1",
                        "fullName": "Drew Quinn",
                        "displayName": "Drew Quinn",
                        "program": "Other",
                        "level": "PGY-2",
                        "role": "Resident",
                        "segments": [{"start": "2026-05-04", "end": "2026-05-05"}],
                        "schoolType": "other",
                        "continuityClinic": "Monday AM",
                        "dayOff": [],
                        "unavailableRanges": [],
                    }
                ],
                "inpatientAssignments": [
                    {
                        "id": "in-2026-05-04-r1",
                        "date": "2026-05-04",
                        "rotatorId": "r1",
                        "role": "Resident",
                        "source": "Manual",
                    }
                ],
                "outpatientSessions": [
                    {
                        "id": "out-2026-05-04-am-r1",
                        "date": "2026-05-04",
                        "period": "AM",
                        "clinic": "",
                        "provider": "",
                        "rotatorId": "r1",
                        "status": "Scheduled",
                    }
                ],
                "attendings": [
                    {
                        "name": "Alder",
                        "recurringClinics": [
                            {"weekday": "Monday", "session": "AM", "clinicName": "General Neuro", "capacity": 1}
                        ],
                        "oneOffDates": [],
                    }
                ],
                "clinicAssignments": [
                    {
                        "id": "clinic-1",
                        "clinicOccurrenceId": "clinic-occurrence::alder::0-monday-AM-general-neuro::2026-05-04::AM",
                        "rotatorId": "r1",
                        "date": "2026-05-04",
                        "session": "AM",
                        "source": "manual",
                    }
                ],
            }
        )
        assert client.post("/api/scheduler/state", json=state).json()["ok"] is True

        response = client.post("/api/scheduler/command", json={"type": "export.word"})
        assert response.status_code == 200
        body = response.json()
        assert body["ok"] is True
        assert body["changed"] is False
        files = body["data"]["files"]
        assert body["message"] == "Built 4 Word exports."
        assert [file["name"] for file in files] == [
            "word-week-schedule.docx",
            "word-week-outpatient-week-grid.docx",
            "word-week-daily-reports.docx",
            "word-week-roster-legend.docx",
        ]
        assert all(file["mimeType"] == "application/vnd.openxmlformats-officedocument.wordprocessingml.document" for file in files)
        decoded = {file["name"]: base64.b64decode(file["base64"]) for file in files}
        assert all(content.startswith(b"PK") for content in decoded.values())
        text_by_name = {file["name"]: _docx_text(file["base64"]) for file in files}
        assert "Word Week - Schedule" in text_by_name["word-week-schedule.docx"]
        assert "Drew Quinn (Resident)" in text_by_name["word-week-schedule.docx"]
        assert "AM General Neuro with Alder: Drew Quinn" in text_by_name["word-week-schedule.docx"]
        assert "Word Week - Daily Reports" in text_by_name["word-week-daily-reports.docx"]
        assert "Daily Team Report" in text_by_name["word-week-daily-reports.docx"]
        assert "Clinic Assignments" in text_by_name["word-week-daily-reports.docx"]
        assert "Word Week - Roster and Legend" in text_by_name["word-week-roster-legend.docx"]
        assert "Word Week - Outpatient Week Grid" in text_by_name["word-week-outpatient-week-grid.docx"]
        assert "Week 1" in text_by_name["word-week-outpatient-week-grid.docx"]
        assert "Monday AM" in text_by_name["word-week-roster-legend.docx"]


def test_rotator_add_update_delete_commands_persist_and_cleanup(tmp_path):
    with TestClient(create_app(data_dir=tmp_path)) as client:
        state = client.get("/api/scheduler/state").json()
        block = {
            "id": "b1",
            "name": "Roster",
            "startDate": "2026-05-04",
            "endDate": "2026-05-08",
            "status": "Draft",
            "generate": {},
            "holidays": [],
        }
        state.update(
            {
                "activeBlockId": "b1",
                "serviceBlocks": [block],
                "rotators": [],
                "inpatientAssignments": [],
                "outpatientSessions": [],
            }
        )
        assert client.post("/api/scheduler/state", json=state).json()["ok"] is True

        add = client.post(
            "/api/scheduler/command",
            json={
                "type": "rotator.add",
                "input": {
                    "fullName": "Drew Quinn",
                    "program": "UT Pediatrics",
                    "level": "PGY-2",
                    "segments": [
                        {"start": "2026-05-04", "end": "2026-05-08", "defaultPhase": "inpatient"}
                    ],
                    "continuityClinic": "Wednesday PM",
                    "dayOff": ["Friday", "Monday", "Nonesday", "Monday"],
                    "unavailableRanges": [
                        {"start": "2026-05-07", "end": "2026-05-07", "label": " Conference "},
                        {"start": "2026-05-15", "end": "2026-05-13", "label": "Backwards"},
                        {"start": "2026-02-31", "end": "2026-02-31", "label": "Impossible"},
                        {"start": "2026-05-08", "end": "2026-05-08", "label": 123},
                    ],
                },
            },
        ).json()
        assert add["ok"] is True
        rotator_id = add["data"]["rotatorId"]
        rotator = add["state"]["rotators"][0]
        assert rotator["id"].startswith("rot-drew-quinn-")
        assert rotator["role"] == "Resident"
        assert rotator["schoolType"] == "ut-peds"
        assert rotator["continuityClinic"] == "Wednesday PM"
        assert rotator["segments"] == [{"start": "2026-05-04", "end": "2026-05-08", "defaultPhase": "inpatient"}]
        assert rotator["dayOff"] == ["Monday", "Friday"]
        assert rotator["unavailableRanges"] == [
            {"start": "2026-05-07", "end": "2026-05-07", "label": "Conference"},
            {"start": "2026-05-08", "end": "2026-05-08"},
        ]

        update = client.post(
            "/api/scheduler/command",
            json={
                "type": "rotator.update",
                "input": {
                    "rotatorId": rotator_id,
                    "patch": {
                        "displayName": "Ari K.",
                        "continuityClinic": "Thursday AM",
                        "segments": [
                            {"start": "2026-05-04", "end": "2026-05-08", "defaultPhase": "inpatient"},
                            {"start": "2026-05-11", "end": "2026-05-15", "defaultPhase": "outpatient"},
                        ],
                        "dayOff": ["Tuesday", "Friday"],
                        "unavailableRanges": [
                            {"start": "2026-05-12", "end": "2026-05-13", "label": "Vacation"},
                            {"start": "2026-05-20", "end": "2026-05-18", "label": "Backwards"},
                            {"start": "not-a-date", "end": "2026-05-18", "label": "Bad"},
                        ],
                    },
                },
            },
        ).json()
        assert update["ok"] is True
        assert update["state"]["rotators"][0]["displayName"] == "Ari K."
        assert update["state"]["rotators"][0]["continuityClinic"] == "Thursday AM"
        assert update["state"]["rotators"][0]["segments"] == [
            {"start": "2026-05-04", "end": "2026-05-08", "defaultPhase": "inpatient"},
            {"start": "2026-05-11", "end": "2026-05-15", "defaultPhase": "outpatient"},
        ]
        assert update["state"]["rotators"][0]["dayOff"] == ["Tuesday", "Friday"]
        assert update["state"]["rotators"][0]["unavailableRanges"] == [
            {"start": "2026-05-12", "end": "2026-05-13", "label": "Vacation"}
        ]

        cleared = client.post(
            "/api/scheduler/command",
            json={
                "type": "rotator.update",
                "input": {
                    "rotatorId": rotator_id,
                    "patch": {
                        "dayOff": [],
                        "unavailableRanges": [],
                    },
                },
            },
        ).json()
        assert cleared["ok"] is True
        assert cleared["state"]["rotators"][0]["dayOff"] == []
        assert cleared["state"]["rotators"][0]["unavailableRanges"] == []

        seeded = cleared["state"]
        seeded["inpatientAssignments"] = [
            {"id": "ip-1", "date": "2026-05-04", "rotatorId": rotator_id, "role": "Resident", "source": "Manual"}
        ]
        seeded["outpatientSessions"] = [
            {
                "id": "op-1",
                "date": "2026-05-05",
                "period": "AM",
                "clinic": "QRS",
                "provider": "",
                "rotatorId": rotator_id,
                "status": "Scheduled",
                "source": "Manual",
            }
        ]
        seeded["clinicAssignments"] = [
            {
                "id": "clinic-1",
                "clinicOccurrenceId": "occ-1",
                "rotatorId": rotator_id,
                "date": "2026-05-05",
                "session": "AM",
                "source": "Manual",
            }
        ]
        assert client.post("/api/scheduler/state", json=seeded).json()["ok"] is True

        delete = client.post(
            "/api/scheduler/command",
            json={"type": "rotator.delete", "input": {"rotatorRef": "Ari K."}},
        ).json()
        assert delete["ok"] is True
        assert delete["state"]["rotators"] == []
        assert delete["state"]["inpatientAssignments"] == []
        assert delete["state"]["outpatientSessions"] == []
        assert delete["state"]["clinicAssignments"] == []

        after = client.get("/api/scheduler/state").json()
        assert after["rotators"] == []
        assert after["inpatientAssignments"] == []
        assert after["outpatientSessions"] == []
        assert after["clinicAssignments"] == []


def test_rotator_add_update_preserves_methodist_rotation_metadata(tmp_path):
    with TestClient(create_app(data_dir=tmp_path)) as client:
        state = client.get("/api/scheduler/state").json()
        state["rotators"] = []
        assert client.post("/api/scheduler/state", json=state).json()["ok"] is True

        add = client.post(
            "/api/scheduler/command",
            json={
                "type": "rotator.add",
                "input": {
                    "fullName": "Maya Lopez",
                    "program": "Methodist",
                    "level": "PGY-3",
                    "segments": [{"start": "2026-06-20", "end": "2026-07-17"}],
                    "rotationStartDate": "2026-06-20",
                    "methodistStartSide": "inpatient",
                },
            },
        ).json()

        assert add["ok"] is True
        rotator = add["state"]["rotators"][0]
        assert rotator["schoolType"] == "methodist"
        assert rotator["rotationStartDate"] == "2026-06-20"
        assert rotator["methodistStartSide"] == "inpatient"

        update = client.post(
            "/api/scheduler/command",
            json={
                "type": "rotator.update",
                "input": {
                    "rotatorId": rotator["id"],
                    "patch": {
                        "rotationStartDate": "2026-06-21",
                        "methodistStartSide": "outpatient",
                    },
                },
            },
        ).json()
        assert update["ok"] is True
        assert update["state"]["rotators"][0]["rotationStartDate"] == "2026-06-21"
        assert update["state"]["rotators"][0]["methodistStartSide"] == "outpatient"

        cleared = client.post(
            "/api/scheduler/command",
            json={
                "type": "rotator.update",
                "input": {
                    "rotatorId": rotator["id"],
                    "patch": {
                        "rotationStartDate": "not-a-date",
                        "methodistStartSide": "sideways",
                    },
                },
            },
        ).json()
        assert cleared["ok"] is True
        assert cleared["state"]["rotators"][0]["rotationStartDate"] == ""
        assert cleared["state"]["rotators"][0]["methodistStartSide"] == ""


def test_roster_dedupe_and_bulk_delete_repoint_or_cleanup_assignments(tmp_path):
    with TestClient(create_app(data_dir=tmp_path)) as client:
        state = client.get("/api/scheduler/state").json()
        state.update(
            {
                "rotators": [
                    {
                        "id": "r1",
                        "fullName": "Drew Quinn",
                        "displayName": "Drew Quinn",
                        "program": "UT Pediatrics",
                        "level": "PGY-2",
                        "role": "Resident",
                        "segments": [{"start": "2026-05-04", "end": "2026-05-08"}],
                        "schoolType": "ut-peds",
                        "continuityClinic": "",
                        "dayOff": [],
                        "unavailableRanges": [],
                    },
                    {
                        "id": "r2",
                        "fullName": "Drew Quinn",
                        "displayName": "Drew Quinn",
                        "program": "UT Pediatrics",
                        "level": "PGY-2",
                        "role": "Resident",
                        "segments": [
                            {"start": "2026-05-04", "end": "2026-05-08", "defaultPhase": "inpatient"},
                            {"start": "2026-05-11", "end": "2026-05-15", "defaultPhase": "outpatient"},
                        ],
                        "schoolType": "ut-peds",
                        "continuityClinic": "",
                        "dayOff": [],
                        "unavailableRanges": [],
                    },
                    {
                        "id": "r3",
                        "fullName": "Blake Lee",
                        "displayName": "Blake Lee",
                        "program": "Other",
                        "level": "PGY-3",
                        "role": "Resident",
                        "segments": [{"start": "2026-05-04", "end": "2026-05-08"}],
                        "schoolType": "other",
                        "continuityClinic": "",
                        "dayOff": [],
                        "unavailableRanges": [],
                    },
                ],
                "inpatientAssignments": [
                    {"id": "ip-1", "date": "2026-05-04", "rotatorId": "r2", "role": "Resident", "source": "Manual"}
                ],
                "outpatientSessions": [
                    {
                        "id": "op-1",
                        "date": "2026-05-05",
                        "period": "AM",
                        "clinic": "QRS",
                        "provider": "",
                        "rotatorId": "r2",
                        "status": "Scheduled",
                        "source": "Manual",
                    }
                ],
                "clinicAssignments": [
                    {
                        "id": "clinic-1",
                        "clinicOccurrenceId": "occ-1",
                        "rotatorId": "r2",
                        "date": "2026-05-05",
                        "session": "AM",
                        "source": "Manual",
                    }
                ],
            }
        )
        assert client.post("/api/scheduler/state", json=state).json()["ok"] is True

        deduped = client.post("/api/scheduler/command", json={"type": "roster.dedupe"}).json()
        assert deduped["ok"] is True
        assert deduped["data"] == {"removedCount": 1}
        assert deduped["message"] == "Removed 1 duplicate rotators."
        assert [rotator["id"] for rotator in deduped["state"]["rotators"]] == ["r1", "r3"]
        ari = deduped["state"]["rotators"][0]
        assert ari["segments"] == [
            {"start": "2026-05-04", "end": "2026-05-08", "defaultPhase": "inpatient"},
            {"start": "2026-05-11", "end": "2026-05-15", "defaultPhase": "outpatient"},
        ]
        assert deduped["state"]["inpatientAssignments"][0]["rotatorId"] == "r1"
        assert deduped["state"]["outpatientSessions"][0]["rotatorId"] == "r1"
        assert deduped["state"]["clinicAssignments"][0]["rotatorId"] == "r1"

        deleted = client.post(
            "/api/scheduler/command",
            json={"type": "rotators.delete", "input": {"rotatorRefs": ["Drew Quinn", "r3"]}},
        ).json()
        assert deleted["ok"] is True
        assert deleted["data"] == {"rotatorIds": ["r1", "r3"]}
        assert deleted["state"]["rotators"] == []
        assert deleted["state"]["inpatientAssignments"] == []
        assert deleted["state"]["outpatientSessions"] == []
        assert deleted["state"]["clinicAssignments"] == []

        after = client.get("/api/scheduler/state").json()
        assert after["rotators"] == []
        assert after["inpatientAssignments"] == []
        assert after["outpatientSessions"] == []
        assert after["clinicAssignments"] == []


def test_unknown_command_returns_ok_false_not_http_error(tmp_path):
    with TestClient(create_app(data_dir=tmp_path)) as client:
        response = client.post("/api/scheduler/command", json={"type": "does.not.exist"})
        assert response.status_code == 200
        body = response.json()
        assert body["ok"] is False
        assert body["error"]["code"] == "unknown_command"
        # No mutation persisted.
        assert client.get("/api/scheduler/state").json()["rules"]["maxConsecutiveInpatientDays"] == 6


def test_malformed_json_body_is_400(tmp_path):
    with TestClient(create_app(data_dir=tmp_path)) as client:
        response = client.post(
            "/api/scheduler/command",
            content=b"{not json",
            headers={"content-type": "application/json"},
        )
        assert response.status_code == 400


def test_non_object_command_is_400(tmp_path):
    with TestClient(create_app(data_dir=tmp_path)) as client:
        response = client.post("/api/scheduler/command", json=["not", "an", "object"])
        assert response.status_code == 400


def test_assign_range_no_op_reports_no_changes(tmp_path):
    # Coordinator 2026-07-29 #5: when every requested day is skipped, the command
    # must say so instead of returning the canned "Assigned ..." success text.
    with TestClient(create_app(data_dir=tmp_path)) as client:
        state = client.get("/api/scheduler/state").json()
        block = {
            "id": "b1",
            "name": "Editable",
            "startDate": "2026-05-04",
            "endDate": "2026-05-08",
            "status": "Draft",
            "generate": {},
            "holidays": [],
        }
        state.update(
            {
                "activeBlockId": "b1",
                "serviceBlocks": [block],
                "rotators": [
                    {
                        "id": "r1",
                        "fullName": "Drew Quinn",
                        "displayName": "Drew Quinn",
                        "program": "Other",
                        "level": "PGY-2",
                        "role": "Resident",
                        "segments": [{"start": "2026-05-04", "end": "2026-05-08"}],
                        "schoolType": "other",
                        "continuityClinic": "",
                        "dayOff": ["Monday"],
                        "unavailableRanges": [],
                    }
                ],
                "inpatientAssignments": [],
                "outpatientSessions": [],
            }
        )
        assert client.post("/api/scheduler/state", json=state).json()["ok"] is True

        response = client.post(
            "/api/scheduler/command",
            json={
                "type": "assign.range",
                "input": {
                    "rotatorRef": "r1",
                    "startDate": "2026-05-04",
                    "endDate": "2026-05-04",
                    "phase": "outpatient",
                },
            },
        )
        assert response.status_code == 200
        body = response.json()
        assert body["ok"] is True
        assert body["changed"] is False
        assert body["message"].startswith("No changes for Drew Quinn")
